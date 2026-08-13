from __future__ import annotations

import copy
import importlib.util
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, PngImagePlugin


SKILL_ROOT = Path(__file__).resolve().parents[1] / "skills" / "legal-network-verification"
SCRIPT_PATH = SKILL_ROOT / "scripts" / "network_workpaper.py"
SPEC = importlib.util.spec_from_file_location("network_workpaper_v20", SCRIPT_PATH)
network_workpaper = importlib.util.module_from_spec(SPEC)
assert SPEC and SPEC.loader
SPEC.loader.exec_module(network_workpaper)


def make_template(path: Path) -> None:
    document = Document()
    document.add_paragraph("关于【项目名称】的网络查询记录")
    document.add_paragraph("查验事项：【查验事项】")
    document.add_paragraph("核查期间：【核查期间】")
    document.add_paragraph("查询时间：【查询时间】")
    document.add_paragraph("查询地点：【查询地点】")
    document.add_paragraph("查询人：【查询人】")
    table = document.add_table(rows=2, cols=3)
    table.cell(0, 0).text = "主体角色"
    table.cell(0, 1).text = "姓名/名称"
    table.cell(0, 2).text = "统一社会信用代码/身份证号码"
    table.cell(1, 0).text = "【SUBJECT_ROWS】"
    document.add_paragraph("【QUERY_RESULT_ITEMS】")
    footer = document.sections[0].footer.paragraphs[0]
    field_run = footer.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    field_run._r.extend((begin, instruction, end))
    document.save(path)


def docx_text(path: Path) -> str:
    document = Document(path)
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        parts.extend(cell.text for row in table.rows for cell in row.cells)
    return "\n".join(parts)


def scope_item() -> dict:
    return {
        "scope_item_id": "SCOPE-20-001",
        "subject_id": "C20",
        "matter_category_id": "company.litigation_judgments",
        "site_id": "official-example",
        "site_name": "示例官方平台",
        "site_basis": "user_specified",
        "allowed_domains": ["example.gov.cn"],
        "query_term_mode": "exact_subject_name",
        "conditions": [
            {"condition_id": "COND-20-001", "field": "period", "value": "2023年至查询日"},
            {"condition_id": "COND-20-002", "field": "region", "value": "全国"},
        ],
    }


def make_nonblank_screenshot(root: Path) -> Path:
    path = root / "01-内部底稿" / "截图" / "NQ-20-001.png"
    path.parent.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (320, 200), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((20, 20, 300, 180), outline="black", width=3)
    draw.text((40, 90), "0 RESULTS", fill="black")
    metadata = PngImagePlugin.PngInfo()
    metadata.add_text("EvidenceId", "NQ-20-001")
    metadata.add_text("QueriedAt", "2026-08-12T10:00:04+08:00")
    metadata.add_text("Subject", "示例企业有限公司")
    metadata.add_text("Source", "示例官方平台")
    metadata.add_text("CaptureKind", "watermarked_page_only")
    metadata.add_text("SourceContentVerified", "heuristic_pre_watermark_gate_passed")
    image.save(path, pnginfo=metadata)
    return path


def base_run(root: Path) -> dict:
    make_nonblank_screenshot(root)
    scope = scope_item()
    execution = {
        "execution_id": "EXEC-20-001",
        "scope_item_id": scope["scope_item_id"],
        "browser_surface": "chrome",
        "browser_selection": "user_explicit",
        "execution_state": "completed",
        "access_state": "authenticated",
        "submit_state": "submitted",
        "block_reason": "",
        "safe_resume_location": "https://sub.example.gov.cn/search",
        "attempts": [
            {"sequence": 1, "at": "2026-08-12T10:00:00+08:00", "event": "access_probe"},
            {"sequence": 2, "at": "2026-08-12T10:00:01+08:00", "event": "page_read"},
            {"sequence": 3, "at": "2026-08-12T10:00:02+08:00", "event": "query_submitted"},
            {"sequence": 4, "at": "2026-08-12T10:00:03+08:00", "event": "result_region_loaded"},
            {"sequence": 5, "at": "2026-08-12T10:00:04+08:00", "event": "screenshot"},
        ],
        "completion_evidence": {
            "query_submitted": True,
            "result_region_loaded": True,
            "result_signal": "explicit_zero_count",
            "observed_at": "2026-08-12T10:00:04+08:00",
            "evidence_id": "NQ-20-001",
            "capture_review": "conditions_and_result_visible",
        },
    }
    query = {
        "evidence_id": "NQ-20-001",
        "execution_id": execution["execution_id"],
        "scope_item_id": scope["scope_item_id"],
        "subject_id": "C20",
        "matter_category_id": scope["matter_category_id"],
        "site_id": scope["site_id"],
        "site_name": scope["site_name"],
        "query_term_mode": scope["query_term_mode"],
        "condition_ids": ["COND-20-001", "COND-20-002"],
        "conditions": copy.deepcopy(scope["conditions"]),
        "url": "https://sub.example.gov.cn/search",
        "query_time": "2026-08-12T10:00:04+08:00",
        "query_terms": "完整主体名称：示例企业有限公司",
        "filters": "period：2023年至查询日；region：全国",
        "status": "no_match_displayed",
        "result_summary": "页面明确显示零条匹配记录。",
        "identity_assessment": "在本次查询条件下未显示可确认匹配的结果。",
        "follow_up": "必要时结合其他书面材料复核。",
        "screenshot_path": "01-内部底稿/截图/NQ-20-001.png",
        "capture_kind": "watermarked_page_only",
    }
    return {
        "schema_version": "1.2",
        "run": {
            "run_id": "NQ-20-TEST",
            "timezone": "Asia/Shanghai",
            "profile": "general-company",
            "project": {
                "name": "匿名网络核查项目",
                "short_name": "匿名项目",
                "matter": "企业公开信息核查",
                "period": "2023年至查询日",
            },
            "formal_record": {
                "query_date": "2026年8月12日",
                "query_location": "北京市",
                "query_people": ["测试律师"],
            },
        },
        "subjects": [
            {
                "subject_id": "C20",
                "type": "company",
                "name": "示例企业有限公司",
                "role": "核查对象",
                "formal_identifier_mode": "user_fill",
            }
        ],
        "query_scope": {
            "status": "user_confirmed",
            "confirmed_at": "2026-08-12T09:59:00+08:00",
            "selection_mode": "custom",
            "items": [scope],
        },
        "query_executions": [execution],
        "queries": [query],
        "opinion_wording_requested": False,
    }


def set_pending_login(run: dict, *, post_submit: bool = False) -> None:
    execution = run["query_executions"][0]
    execution.update(
        {
            "execution_state": "pending_user_action",
            "access_state": "session_expired" if post_submit else "login_required",
            "submit_state": "unknown" if post_submit else "not_submitted",
            "block_reason": "submit_state_unknown" if post_submit else "login_required",
            "safe_resume_location": "https://sub.example.gov.cn/login",
            "completion_evidence": {
                "query_submitted": post_submit,
                "result_region_loaded": False,
                "result_signal": "not_observed",
                "observed_at": "",
                "evidence_id": "",
                "capture_review": "not_reviewed",
            },
        }
    )
    if post_submit:
        execution["attempts"] = [
            {"sequence": 1, "at": "2026-08-12T10:00:00+08:00", "event": "access_probe"},
            {"sequence": 2, "at": "2026-08-12T10:00:01+08:00", "event": "query_submitted"},
            {"sequence": 3, "at": "2026-08-12T10:00:02+08:00", "event": "session_expired"},
        ]
    else:
        execution["attempts"] = [
            {"sequence": 1, "at": "2026-08-12T10:00:00+08:00", "event": "access_probe"},
            {"sequence": 2, "at": "2026-08-12T10:00:01+08:00", "event": "login_required"},
        ]
    run["queries"] = []


def set_blocked(run: dict, reason: str, event: str) -> None:
    execution = run["query_executions"][0]
    execution.update(
        {
            "execution_state": "blocked",
            "access_state": "limited",
            "submit_state": "not_submitted",
            "block_reason": reason,
            "attempts": [
                {"sequence": 1, "at": "2026-08-12T10:00:00+08:00", "event": "access_probe"},
                {"sequence": 2, "at": "2026-08-12T10:00:01+08:00", "event": event},
            ],
            "completion_evidence": {
                "query_submitted": False,
                "result_region_loaded": False,
                "result_signal": "not_observed",
                "observed_at": "",
                "evidence_id": "",
                "capture_review": "not_reviewed",
            },
        }
    )
    query = run["queries"][0]
    query.update(
        {
            "status": "access_limited",
            "result_summary": network_workpaper.CANONICAL_RESTRICTED_RESULT,
            "identity_assessment": network_workpaper.CANONICAL_RESTRICTED_ASSESSMENT,
            "follow_up": network_workpaper.CANONICAL_RESTRICTED_FOLLOW_UP,
            "screenshot_path": "",
            "capture_kind": "",
        }
    )


class SchemaV20ExecutionTests(unittest.TestCase):
    def test_completed_zero_result_passes_strict_gate_and_completion_preview(self) -> None:
        with tempfile.TemporaryDirectory() as temp_text:
            root = Path(temp_text)
            run = base_run(root)
            network_workpaper.validate_run(run, workpaper_root=root, formal_mode="final")
            preview = network_workpaper.completion_preview(run, workpaper_root=root)
            self.assertTrue(preview["final_ready"])
            self.assertEqual(preview["counts"]["completed"], 1)
            self.assertEqual(preview["blocking_items"], [])
            self.assertTrue(network_workpaper.scope_preview(run)["query_allowed"])

    def test_prepare_defaults_to_12_and_legacy_versions_are_read_only(self) -> None:
        prepared = network_workpaper.prepare_run(
            {
                "project": {
                    "name": "匿名项目",
                    "short_name": "匿名",
                    "matter": "公开核查",
                    "period": "截至查询日",
                },
                "subjects": [
                    {
                        "subject_id": "C20",
                        "type": "company",
                        "name": "示例企业有限公司",
                        "role": "核查对象",
                    }
                ],
            }
        )
        self.assertEqual(prepared["schema_version"], "1.2")
        self.assertEqual(prepared["query_executions"], [])
        for version in ("1.0", "1.1"):
            with self.subTest(version=version):
                with self.assertRaises(network_workpaper.ValidationError):
                    network_workpaper.prepare_run(
                        {
                            "schema_version": version,
                            "project": prepared["run"]["project"],
                            "subjects": prepared["subjects"],
                        }
                    )

    def test_missing_execution_builds_draft_completion_views_but_blocks_final(self) -> None:
        with tempfile.TemporaryDirectory() as temp_text:
            root = Path(temp_text)
            run = base_run(root)
            run["query_executions"] = []
            run["queries"] = []
            template = root / "template.docx"
            make_template(template)
            outputs = network_workpaper.build_outputs(
                run,
                workpaper_root=root,
                output_dir=root,
                template_docx=template,
                formal_mode="draft",
                overwrite=False,
            )
            markdown = next(path for path in outputs if path.suffix == ".md").read_text(encoding="utf-8")
            formal = docx_text(next(path for path in outputs if path.suffix == ".docx"))
            self.assertIn("## 一、范围项完成度", markdown)
            self.assertIn("未开始", markdown)
            self.assertIn("本项尚未开始", formal)
            self.assertNotIn("not_started", formal)
            preview = network_workpaper.completion_preview(run, workpaper_root=root)
            self.assertFalse(preview["final_ready"])
            with self.assertRaises(network_workpaper.ValidationError):
                network_workpaper.validate_run(run, workpaper_root=root, formal_mode="final")

    def test_pending_login_is_recoverable_in_draft_and_never_a_zero_result(self) -> None:
        with tempfile.TemporaryDirectory() as temp_text:
            root = Path(temp_text)
            run = base_run(root)
            set_pending_login(run)
            network_workpaper.validate_run(run, workpaper_root=root, formal_mode="draft")
            template = root / "template.docx"
            make_template(template)
            outputs = network_workpaper.build_outputs(
                run,
                workpaper_root=root,
                output_dir=root,
                template_docx=template,
                formal_mode="draft",
                overwrite=False,
            )
            markdown = next(path for path in outputs if path.suffix == ".md").read_text(encoding="utf-8")
            formal = docx_text(next(path for path in outputs if path.suffix == ".docx"))
            self.assertIn("待登录", markdown)
            self.assertIn("待用户完成登录", formal)
            self.assertNotIn("未显示可确认匹配", formal)
            with self.assertRaises(network_workpaper.ValidationError):
                network_workpaper.validate_run(run, workpaper_root=root, formal_mode="final")

    def test_handoff_forbids_all_page_operations_until_user_resume(self) -> None:
        with tempfile.TemporaryDirectory() as temp_text:
            root = Path(temp_text)
            for event in (
                "access_probe",
                "browser_fallback",
                "page_read",
                "click",
                "screenshot",
                "query_submitted",
                "result_region_loaded",
            ):
                run = base_run(root)
                set_pending_login(run)
                run["query_executions"][0]["attempts"].append(
                    {"sequence": 3, "at": "2026-08-12T10:00:02+08:00", "event": event}
                )
                with self.subTest(event=event):
                    with self.assertRaises(network_workpaper.ValidationError):
                        network_workpaper.validate_run(run, workpaper_root=root)

    def test_manual_handoff_can_resume_same_execution_after_explicit_user_confirmation(self) -> None:
        with tempfile.TemporaryDirectory() as temp_text:
            root = Path(temp_text)
            run = base_run(root)
            execution = run["query_executions"][0]
            execution["attempts"] = [
                {"sequence": 1, "at": "2026-08-12T10:00:00+08:00", "event": "access_probe"},
                {"sequence": 2, "at": "2026-08-12T10:00:01+08:00", "event": "login_required"},
                {"sequence": 3, "at": "2026-08-12T10:00:02+08:00", "event": "user_resume_confirmed"},
                {"sequence": 4, "at": "2026-08-12T10:00:02+08:00", "event": "resume_safety_check_passed"},
                {"sequence": 5, "at": "2026-08-12T10:00:03+08:00", "event": "page_read"},
                {"sequence": 6, "at": "2026-08-12T10:00:04+08:00", "event": "query_submitted"},
                {"sequence": 7, "at": "2026-08-12T10:00:05+08:00", "event": "result_region_loaded"},
                {"sequence": 8, "at": "2026-08-12T10:00:06+08:00", "event": "screenshot"},
            ]
            execution["completion_evidence"]["observed_at"] = "2026-08-12T10:00:06+08:00"
            run["queries"][0]["query_time"] = "2026-08-12T10:00:06+08:00"
            make_nonblank_screenshot(root)
            screenshot_path = root / run["queries"][0]["screenshot_path"]
            with Image.open(screenshot_path) as opened:
                pixels = opened.convert("RGB")
            metadata = PngImagePlugin.PngInfo()
            for key, value in {
                "EvidenceId": "NQ-20-001",
                "QueriedAt": "2026-08-12T10:00:06+08:00",
                "Subject": "示例企业有限公司",
                "Source": "示例官方平台",
                "CaptureKind": "watermarked_page_only",
                "SourceContentVerified": "heuristic_pre_watermark_gate_passed",
            }.items():
                metadata.add_text(key, value)
            pixels.save(screenshot_path, pnginfo=metadata)
            network_workpaper.validate_run(run, workpaper_root=root, formal_mode="final")

    def test_human_verification_is_shown_as_pending_not_as_a_result(self) -> None:
        with tempfile.TemporaryDirectory() as temp_text:
            root = Path(temp_text)
            run = base_run(root)
            set_pending_login(run)
            execution = run["query_executions"][0]
            execution["access_state"] = "human_verification_required"
            execution["block_reason"] = "human_verification_required"
            execution["attempts"][1]["event"] = "human_verification_required"
            markdown = network_workpaper.build_internal_markdown(
                run, run["subjects"][0], run["queries"]
            )
            self.assertIn("待人工验证", markdown)
            network_workpaper.validate_run(run, workpaper_root=root, formal_mode="draft")

    def test_explicit_browser_cannot_fallback_and_automatic_fallback_occurs_once(self) -> None:
        with tempfile.TemporaryDirectory() as temp_text:
            root = Path(temp_text)
            run = base_run(root)
            run["query_executions"][0]["attempts"].insert(
                1, {"sequence": 2, "at": "2026-08-12T10:00:00+08:00", "event": "browser_fallback"}
            )
            for index, attempt in enumerate(run["query_executions"][0]["attempts"], start=1):
                attempt["sequence"] = index
            with self.assertRaises(network_workpaper.ValidationError):
                network_workpaper.validate_run(run, workpaper_root=root)

            run["query_executions"][0]["browser_selection"] = "automatic_fallback"
            network_workpaper.validate_run(run, workpaper_root=root)
            duplicate = copy.deepcopy(run)
            duplicate["query_executions"][0]["attempts"].insert(
                2, {"sequence": 3, "at": "2026-08-12T10:00:00+08:00", "event": "browser_fallback"}
            )
            for index, attempt in enumerate(duplicate["query_executions"][0]["attempts"], start=1):
                attempt["sequence"] = index
            with self.assertRaises(network_workpaper.ValidationError):
                network_workpaper.validate_run(duplicate, workpaper_root=root)

    def test_post_submit_challenge_requires_unknown_state_and_explicit_resubmit(self) -> None:
        with tempfile.TemporaryDirectory() as temp_text:
            root = Path(temp_text)
            run = base_run(root)
            set_pending_login(run, post_submit=True)
            network_workpaper.validate_run(run, workpaper_root=root, formal_mode="draft")
            with self.assertRaises(network_workpaper.ValidationError):
                network_workpaper.validate_run(run, workpaper_root=root, formal_mode="final")

            invalid = copy.deepcopy(run)
            execution = invalid["query_executions"][0]
            execution["execution_state"] = "running"
            execution["access_state"] = "authenticated"
            execution["submit_state"] = "submitted"
            execution["block_reason"] = ""
            execution["attempts"].extend(
                [
                    {"sequence": 4, "at": "2026-08-12T10:00:03+08:00", "event": "user_resume_confirmed"},
                    {"sequence": 5, "at": "2026-08-12T10:00:04+08:00", "event": "query_submitted"},
                ]
            )
            with self.assertRaises(network_workpaper.ValidationError):
                network_workpaper.validate_run(invalid, workpaper_root=root)

            recovered = base_run(root)
            execution = recovered["query_executions"][0]
            execution["attempts"] = [
                {"sequence": 1, "at": "2026-08-12T10:00:00+08:00", "event": "query_submitted"},
                {"sequence": 2, "at": "2026-08-12T10:00:01+08:00", "event": "result_region_loaded"},
                {"sequence": 3, "at": "2026-08-12T10:00:02+08:00", "event": "session_expired"},
                {"sequence": 4, "at": "2026-08-12T10:00:03+08:00", "event": "user_resume_confirmed"},
                {"sequence": 5, "at": "2026-08-12T10:00:03+08:00", "event": "resume_safety_check_passed"},
                {"sequence": 6, "at": "2026-08-12T10:00:04+08:00", "event": "explicit_resubmit_confirmed"},
                {"sequence": 7, "at": "2026-08-12T10:00:05+08:00", "event": "query_submitted"},
                {"sequence": 8, "at": "2026-08-12T10:00:06+08:00", "event": "result_region_loaded"},
                {"sequence": 9, "at": "2026-08-12T10:00:07+08:00", "event": "screenshot"},
            ]
            execution["completion_evidence"]["observed_at"] = "2026-08-12T10:00:07+08:00"
            recovered["queries"][0]["query_time"] = "2026-08-12T10:00:07+08:00"
            screenshot_path = make_nonblank_screenshot(root)
            with Image.open(screenshot_path) as opened:
                pixels = opened.convert("RGB")
            metadata = PngImagePlugin.PngInfo()
            for key, value in {
                "EvidenceId": "NQ-20-001",
                "QueriedAt": "2026-08-12T10:00:07+08:00",
                "Subject": "示例企业有限公司",
                "Source": "示例官方平台",
                "CaptureKind": "watermarked_page_only",
                "SourceContentVerified": "heuristic_pre_watermark_gate_passed",
            }.items():
                metadata.add_text(key, value)
            pixels.save(screenshot_path, pnginfo=metadata)
            network_workpaper.validate_run(recovered, workpaper_root=root, formal_mode="final")

            stale_loaded = copy.deepcopy(recovered)
            stale_loaded["query_executions"][0]["attempts"].pop(7)
            for index, attempt in enumerate(stale_loaded["query_executions"][0]["attempts"], start=1):
                attempt["sequence"] = index
            with self.assertRaises(network_workpaper.ValidationError):
                network_workpaper.validate_run(stale_loaded, workpaper_root=root)

    def test_forbidden_credential_or_browser_state_events_and_fields_fail_closed(self) -> None:
        with tempfile.TemporaryDirectory() as temp_text:
            root = Path(temp_text)
            for event in ("password_read", "cookie_read", "storage_read", "token_read", "profile_read"):
                run = base_run(root)
                run["query_executions"][0]["attempts"][0]["event"] = event
                with self.subTest(event=event):
                    with self.assertRaises(network_workpaper.ValidationError):
                        network_workpaper.validate_run(run, workpaper_root=root)
            run = base_run(root)
            run["query_executions"][0]["profile_name"] = "research"
            with self.assertRaises(network_workpaper.ValidationError):
                network_workpaper.validate_run(run, workpaper_root=root)

    def test_blank_image_full_app_capture_and_login_zero_cannot_prove_no_match(self) -> None:
        with tempfile.TemporaryDirectory() as temp_text:
            root = Path(temp_text)
            run = base_run(root)
            Image.new("RGB", (320, 200), "white").save(
                root / run["queries"][0]["screenshot_path"]
            )
            with self.assertRaises(network_workpaper.ValidationError):
                network_workpaper.validate_run(run, workpaper_root=root)

            run = base_run(root)
            run["queries"][0]["capture_kind"] = "full_app_screenshot"
            with self.assertRaises(network_workpaper.ValidationError):
                network_workpaper.validate_run(run, workpaper_root=root)

            run = base_run(root)
            set_pending_login(run)
            login_result = copy.deepcopy(base_run(root)["queries"][0])
            login_result["result_summary"] = "登录页面显示数字0。"
            run["queries"] = [login_result]
            with self.assertRaises(network_workpaper.ValidationError):
                network_workpaper.validate_run(run, workpaper_root=root)

    def test_safe_resume_location_rejects_userinfo_query_and_fragment(self) -> None:
        with tempfile.TemporaryDirectory() as temp_text:
            root = Path(temp_text)
            for location in (
                "https://user@sub.example.gov.cn/login",
                "https://sub.example.gov.cn/login?next=search",
                "https://sub.example.gov.cn/login#form",
                "https://evil-example.gov.cn/login",
            ):
                run = base_run(root)
                set_pending_login(run)
                run["query_executions"][0]["safe_resume_location"] = location
                with self.subTest(location=location):
                    with self.assertRaises(network_workpaper.ValidationError):
                        network_workpaper.validate_run(run, workpaper_root=root)

    def test_access_blocks_and_failures_map_one_to_one_without_negative_claims(self) -> None:
        with tempfile.TemporaryDirectory() as temp_text:
            root = Path(temp_text)
            for reason, event in (
                ("rate_limited", "rate_limited"),
                ("permission_required", "permission_required"),
                ("paywall", "paywall"),
                ("site_maintenance", "site_maintenance"),
            ):
                run = base_run(root)
                set_blocked(run, reason, event)
                with self.subTest(reason=reason):
                    network_workpaper.validate_run(run, workpaper_root=root, formal_mode="final")
                    text = network_workpaper.concise_formal_query_text(1, run["queries"][0], "示例企业有限公司")
                    self.assertIn("暂不作是否存在负面记录的结论", text)
                    self.assertNotIn(reason, text)

            run = base_run(root)
            set_blocked(run, "network_error", "network_error")
            run["query_executions"][0]["execution_state"] = "failed"
            run["queries"][0]["status"] = "failed"
            run["queries"][0]["result_summary"] = network_workpaper.CANONICAL_FAILED_RESULT
            network_workpaper.validate_run(run, workpaper_root=root, formal_mode="final")
            wrong = copy.deepcopy(run)
            wrong["queries"][0]["status"] = "access_limited"
            with self.assertRaises(network_workpaper.ValidationError):
                network_workpaper.validate_run(wrong, workpaper_root=root)

    def test_user_stopped_is_terminal_and_allows_truthful_access_limited_final(self) -> None:
        with tempfile.TemporaryDirectory() as temp_text:
            root = Path(temp_text)
            run = base_run(root)
            set_blocked(run, "user_stopped", "login_required")
            execution = run["query_executions"][0]
            execution["access_state"] = "login_required"
            execution["attempts"].append(
                {"sequence": 3, "at": "2026-08-12T10:00:02+08:00", "event": "user_stopped"}
            )
            network_workpaper.validate_run(run, workpaper_root=root, formal_mode="final")
            invalid = copy.deepcopy(run)
            invalid["query_executions"][0]["attempts"].append(
                {"sequence": 4, "at": "2026-08-12T10:00:03+08:00", "event": "page_read"}
            )
            with self.assertRaises(network_workpaper.ValidationError):
                network_workpaper.validate_run(invalid, workpaper_root=root)

    def test_sensitive_no_capture_exception_remains_narrow(self) -> None:
        with tempfile.TemporaryDirectory() as temp_text:
            root = Path(temp_text)
            run = base_run(root)
            run["subjects"][0]["type"] = "natural_person"
            run["subjects"][0]["name"] = "示例自然人"
            run["query_scope"]["items"][0]["query_term_mode"] = "user_manual_full_id"
            query = run["queries"][0]
            query["query_term_mode"] = "user_manual_full_id"
            query["query_terms"] = network_workpaper.canonical_query_terms(
                run["subjects"][0], "user_manual_full_id"
            )
            query["screenshot_path"] = ""
            query["capture_kind"] = "not_retained_sensitive_page"
            query["follow_up"] = "因敏感信息保护未留存该页面截图。"
            run["query_executions"][0]["completion_evidence"]["capture_review"] = "sensitive_no_capture"
            run["query_executions"][0]["attempts"][-1] = {
                "sequence": 5,
                "at": "2026-08-12T10:00:04+08:00",
                "event": "sensitive_result_review_confirmed",
            }
            network_workpaper.validate_run(run, workpaper_root=root)
            run["queries"][0]["follow_up"] = "未保留截图。"
            with self.assertRaises(network_workpaper.ValidationError):
                network_workpaper.validate_run(run, workpaper_root=root)

    def test_terminal_result_mapping_and_completion_evidence_are_exactly_one_to_one(self) -> None:
        with tempfile.TemporaryDirectory() as temp_text:
            root = Path(temp_text)
            run = base_run(root)
            run["query_executions"][0]["completion_evidence"]["evidence_id"] = "NQ-OTHER"
            with self.assertRaises(network_workpaper.ValidationError):
                network_workpaper.validate_run(run, workpaper_root=root)

            run = base_run(root)
            duplicate = copy.deepcopy(run["queries"][0])
            duplicate["evidence_id"] = "NQ-20-002"
            run["queries"].append(duplicate)
            with self.assertRaises(network_workpaper.ValidationError):
                network_workpaper.validate_run(run, workpaper_root=root)


if __name__ == "__main__":
    unittest.main()
