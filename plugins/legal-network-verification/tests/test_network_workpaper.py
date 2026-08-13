from __future__ import annotations

import copy
import importlib.util
import json
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from PIL import Image, ImageDraw, PngImagePlugin


PLUGIN_ROOT = Path(__file__).resolve().parents[1]
SKILL_ROOT = PLUGIN_ROOT / "skills" / "legal-network-verification"
MODULE_PATH = SKILL_ROOT / "scripts" / "network_workpaper.py"
SPEC = importlib.util.spec_from_file_location("network_workpaper", MODULE_PATH)
network_workpaper = importlib.util.module_from_spec(SPEC)
assert SPEC.loader is not None
SPEC.loader.exec_module(network_workpaper)

WATERMARK_MODULE_PATH = SKILL_ROOT / "scripts" / "watermark_capture.py"
WATERMARK_SPEC = importlib.util.spec_from_file_location(
    "watermark_capture", WATERMARK_MODULE_PATH
)
watermark_capture = importlib.util.module_from_spec(WATERMARK_SPEC)
assert WATERMARK_SPEC.loader is not None
WATERMARK_SPEC.loader.exec_module(watermark_capture)

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
NS = {"w": W_NS, "r": R_NS, "pr": PKG_REL_NS}
BUNDLED_TEMPLATE = SKILL_ROOT / "assets" / "generic-network-query-record.docx"
SYNTHETIC_FULL_ID = "1101011990" + "03071234"
SYNTHETIC_LEGACY_ID = "1101019003" + "07123"
SYNTHETIC_VALID_CREDIT_CODE = "911100001234567810"
SYNTHETIC_OTHER_VALID_CREDIT_CODE = "911100001234567823"


def make_template(path: Path, *, trailing_subject_row: bool = False) -> None:
    document = Document()
    document.add_heading("网络查询记录", level=1)
    document.add_paragraph("项目名称：【项目名称】")
    document.add_paragraph("查验事项：【查验事项】")
    document.add_paragraph("核查期间：【核查期间】")
    document.add_paragraph("查询时间：【查询时间】")
    document.add_paragraph("查询地点：【查询地点】")
    document.add_paragraph("查询人：【查询人】")
    table = document.add_table(rows=3 if trailing_subject_row else 2, cols=3)
    table.cell(0, 0).text = "主体角色"
    table.cell(0, 1).text = "查询对象"
    table.cell(0, 2).text = "身份号码"
    table.cell(1, 0).text = "【SUBJECT_ROWS】"
    if trailing_subject_row:
        table.cell(2, 0).text = "固定说明行"
    marker = document.add_paragraph("【QUERY_RESULT_ITEMS】")
    marker.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    document.add_paragraph("附件：查询网址截图底稿另行保存。")
    document.add_paragraph("查询律师签署：")
    footer = document.sections[0].footer.paragraphs[0]
    field_run = footer.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    field_run._r.extend((begin, instruction, end))
    document.save(path)


def anonymous_data() -> dict:
    return {
        "schema_version": "1.0",
        "run": {
            "run_id": "NQ-20260717-TEST",
            "timezone": "Asia/Shanghai",
            "profile": "custom",
            "project": {
                "name": "匿名网络核查测试项目",
                "short_name": "匿名测试",
                "matter": "公开网络核查",
                "period": "2023年1月1日至查询日",
            },
            "formal_record": {
                "query_date": "2026年7月17日",
                "query_location": "北京市朝阳区",
                "query_people": ["测试律师"],
            },
        },
        "subjects": [
            {
                "subject_id": "SUB-001",
                "type": "natural_person",
                "name": "甲某",
                "role": "自然人核查对象",
                "associated_entity": "示例科技有限公司",
                "masked_id_number": "1101011990********",
            },
            {
                "subject_id": "SUB-002",
                "type": "company",
                "name": "示例科技有限公司",
                "role": "企业核查对象",
                "associated_entity": "",
                "credit_code": SYNTHETIC_VALID_CREDIT_CODE,
            },
        ],
        "queries": [
            {
                "evidence_id": "NQ-01-01-01",
                "subject_id": "SUB-001",
                "site_id": "court",
                "site_name": "示例裁判文书公开网站",
                "url": "https://example.invalid/case/search?name=%E7%94%B2%E6%9F%90&from=2023-01-01",
                "query_time": "2026-07-17T09:10:11+08:00",
                "query_terms": "姓名及关联机构",
                "filters": "2023年1月1日以后公开记录",
                "status": "identity_match",
                "result_summary": "页面显示1条与姓名及关联机构相符的公开记录",
                "identity_assessment": "姓名和关联机构两项身份要素一致",
                "screenshot_path": "01-内部底稿/截图/NQ-01-01-01.png",
                "capture_kind": "watermarked_page_only",
                "follow_up": "复核记录详情",
            },
            {
                "evidence_id": "NQ-01-02-01",
                "subject_id": "SUB-001",
                "site_id": "enforcement",
                "site_name": "示例执行信息公开网站",
                "url": "https://example.invalid/enforcement/list/search/result/detail/path.with.long.segment?keyword=%E7%94%B2%E6%9F%90&page=1&size=20",
                "query_time": "2026-07-17T09:20:21+08:00",
                "query_terms": "姓名及本案材料所载身份要素",
                "filters": "全部公开栏目",
                "status": "no_match_displayed",
                "result_summary": "页面显示检索结果0条，未显示可确认匹配记录",
                "identity_assessment": "在本次查询条件下未显示可确认匹配的结果",
                "screenshot_path": "01-内部底稿/截图/NQ-01-02-01.png",
                "capture_kind": "watermarked_page_only",
                "follow_up": "",
            },
            {
                "evidence_id": "NQ-01-03-01",
                "subject_id": "SUB-001",
                "site_id": "same-name",
                "site_name": "示例公开检索网站",
                "url": "https://example.invalid/public/search?query=%E7%94%B2%E6%9F%90&category=all",
                "query_time": "2026-07-17T09:30:31+08:00",
                "query_terms": "姓名",
                "filters": "2023年1月1日以后",
                "status": "same_name_candidates",
                "result_summary": "页面显示3条姓名相同的候选记录",
                "identity_assessment": "公开页面信息不足以确认候选记录归属于核查对象",
                "screenshot_path": "01-内部底稿/截图/NQ-01-03-01.png",
                "capture_kind": "watermarked_page_only",
                "follow_up": "结合其他身份要素进一步核对",
            },
            {
                "evidence_id": "NQ-01-04-01",
                "subject_id": "SUB-001",
                "site_id": "company-only",
                "site_name": "示例企业信用网站",
                "url": "https://example.invalid/company/search?key=%E7%94%B2%E6%9F%90",
                "query_time": "2026-07-17T09:40:41+08:00",
                "query_terms": "姓名",
                "filters": "企业名称查询入口",
                "status": "not_applicable",
                "result_summary": "该入口仅接受企业名称或统一社会信用代码",
                "identity_assessment": "该入口不适用于普通自然人姓名核查",
                "screenshot_path": "01-内部底稿/截图/NQ-01-04-01.png",
                "capture_kind": "watermarked_page_only",
                "follow_up": "",
            },
            {
                "evidence_id": "NQ-02-01-01",
                "subject_id": "SUB-002",
                "site_id": "limited",
                "site_name": "示例监管查询网站",
                "url": "https://example.invalid/regulator/search?company=%E7%A4%BA%E4%BE%8B%E7%A7%91%E6%8A%80",
                "query_time": "2026-07-17T09:50:51+08:00",
                "query_terms": "企业名称",
                "filters": "公开查询入口",
                "status": "access_limited",
                "result_summary": "网站触发访问频率限制，未能完整加载查询结果",
                "identity_assessment": "本项查询未完整完成",
                "screenshot_path": "",
                "capture_kind": "",
                "follow_up": "恢复访问后补充查询",
            },
            {
                "evidence_id": "NQ-02-02-01",
                "subject_id": "SUB-002",
                "site_id": "failed",
                "site_name": "示例公告查询网站",
                "url": "https://example.invalid/notice/search?company=%E7%A4%BA%E4%BE%8B%E7%A7%91%E6%8A%80",
                "query_time": "2026-07-17T10:00:01+08:00",
                "query_terms": "企业名称",
                "filters": "公告全文",
                "status": "failed",
                "result_summary": "网站返回错误页面，本次查询未完成",
                "identity_assessment": "无法形成查询结果",
                "screenshot_path": "",
                "capture_kind": "",
                "follow_up": "另行补充查询",
            },
        ],
        "opinion_wording_requested": False,
    }


class NetworkWorkpaperV121Tests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.template = self.root / "template.docx"
        make_template(self.template)
        screenshot_dir = self.root / "01-内部底稿" / "截图"
        screenshot_dir.mkdir(parents=True)
        for evidence_id in ("NQ-01-01-01", "NQ-01-02-01", "NQ-01-03-01", "NQ-01-04-01"):
            Image.new("RGB", (80, 50), "white").save(
                screenshot_dir / f"{evidence_id}.png", format="PNG"
            )
        self.data = anonymous_data()

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def test_two_layer_build_and_docx_contract(self) -> None:
        paths = network_workpaper.build_outputs(
            self.data,
            workpaper_root=self.root,
            output_dir=self.root,
            template_docx=self.template,
            formal_mode="draft",
            layout="two-layer",
            overwrite=False,
        )
        markdown_paths = sorted(path for path in paths if path.suffix == ".md")
        docx_path = next(path for path in paths if path.suffix == ".docx")
        self.assertTrue(all(path.parent.name == "01-内部底稿" for path in markdown_paths))
        self.assertEqual(docx_path.parent.name, "02-正式记录")
        self.assertTrue(docx_path.name.endswith("_草稿.docx"))

        person_markdown = next(path for path in markdown_paths if path.name.startswith("甲某"))
        markdown_text = person_markdown.read_text(encoding="utf-8")
        self.assertIn("(截图/NQ-01-01-01.png)", markdown_text)
        initial_conclusion = markdown_text.split("## 四、初步结论", 1)[1].split(
            "## 五、建议补核事项", 1
        )[0]
        self.assertNotIn("身份判断为", initial_conclusion)
        self.assertNotIn("结论级别为", initial_conclusion)
        self.assertNotIn("公开页面信息不足以确认候选记录归属于核查对象", initial_conclusion)

        with zipfile.ZipFile(docx_path) as archive:
            document_xml = etree.fromstring(archive.read("word/document.xml"))
            rels_xml = etree.fromstring(archive.read("word/_rels/document.xml.rels"))

        formal_text = "".join(document_xml.xpath(".//w:t/text()", namespaces=NS))
        for forbidden in ("身份判断", "结论级别", "NQ-", "复核人"):
            self.assertNotIn(forbidden, formal_text)
        self.assertIn("经查询示例执行信息公开网站", formal_text)
        self.assertIn("仅反映本次查询条件及该网站公开可查询范围", formal_text)

        result_paragraphs = document_xml.xpath(
            ".//w:p[.//w:t[contains(., '就') and contains(., '查询')]]",
            namespaces=NS,
        )
        self.assertEqual(len(result_paragraphs), len(self.data["queries"]))
        for paragraph in result_paragraphs:
            self.assertEqual(paragraph.xpath("string(./w:pPr/w:jc/@w:val)", namespaces=NS), "both")
            self.assertEqual(
                paragraph.xpath("string(./w:pPr/w:wordWrap/@w:val)", namespaces=NS),
                "1",
            )

        relationships = {
            item.get("Id"): item.get("Target")
            for item in rels_xml.xpath("./pr:Relationship", namespaces=NS)
            if item.get("Type", "").endswith("/hyperlink")
        }
        hyperlink_targets = []
        hyperlink_displays = []
        for hyperlink in document_xml.xpath(".//w:hyperlink", namespaces=NS):
            hyperlink_targets.append(relationships[hyperlink.get(f"{{{R_NS}}}id")])
            hyperlink_displays.append(
                "".join(hyperlink.xpath(".//w:t/text()", namespaces=NS))
            )
        expected_urls = [query["url"] for query in self.data["queries"]]
        self.assertCountEqual(hyperlink_targets, expected_urls)
        self.assertCountEqual(
            [display.replace(network_workpaper.ZERO_WIDTH_SPACE, "") for display in hyperlink_displays],
            expected_urls,
        )
        self.assertTrue(all(network_workpaper.ZERO_WIDTH_SPACE in item for item in hyperlink_displays))

    def test_flat_layout_remains_available(self) -> None:
        flat_root = self.root / "flat"
        flat_root.mkdir()
        data = copy.deepcopy(self.data)
        for query in data["queries"]:
            screenshot = query.get("screenshot_path")
            if screenshot:
                source = self.root / screenshot
                destination = flat_root / screenshot
                destination.parent.mkdir(parents=True, exist_ok=True)
                destination.write_bytes(source.read_bytes())
        paths = network_workpaper.build_outputs(
            data,
            workpaper_root=flat_root,
            output_dir=flat_root,
            template_docx=self.template,
            formal_mode="draft",
            layout="flat",
            overwrite=False,
        )
        self.assertTrue(all(path.parent == flat_root for path in paths))
        self.assertTrue(any(path.name.endswith("_草稿.docx") for path in paths))

    def test_final_mode_requires_location_and_people_but_allows_user_fill_identifiers(self) -> None:
        cases = [
            ("query_location", ""),
            ("query_people", []),
        ]
        for field, value in cases:
            with self.subTest(field=field):
                data = copy.deepcopy(self.data)
                data["run"]["formal_record"][field] = value
                with self.assertRaises(network_workpaper.ValidationError):
                    network_workpaper.validate_run(
                        data,
                        workpaper_root=self.root,
                        formal_mode="final",
                    )

        data = copy.deepcopy(self.data)
        data["subjects"][0]["masked_id_number"] = ""
        network_workpaper.validate_run(data, workpaper_root=self.root, formal_mode="final")
        self.assertEqual(
            network_workpaper.manual_identifier_fields_pending(data),
            ["SUB-001.formal_identifier_user_fill", "SUB-002.formal_identifier_user_fill"],
        )
        paths = network_workpaper.build_outputs(
            data, workpaper_root=self.root, output_dir=self.root / "user-fill-final",
            template_docx=self.template, formal_mode="final", layout="two-layer", overwrite=False,
        )
        table = Document(next(path for path in paths if path.suffix == ".docx")).tables[0]
        self.assertEqual([row.cells[2].text for row in table.rows[1:3]], ["", ""])
        self.assertTrue(
            network_workpaper.artifact_audit(
                next(path for path in paths if path.suffix == ".docx")
            )["ok"]
        )

    def test_restricted_and_failed_queries_cannot_claim_no_record(self) -> None:
        for status in ("access_limited", "failed"):
            with self.subTest(status=status):
                data = copy.deepcopy(self.data)
                query = next(item for item in data["queries"] if item["status"] == status)
                query["result_summary"] = "页面未发现记录"
                with self.assertRaises(network_workpaper.ValidationError):
                    network_workpaper.validate_run(
                        data,
                        workpaper_root=self.root,
                        formal_mode="draft",
                    )

    def test_masked_id_contract_and_sensitive_values(self) -> None:
        for invalid in (
            "110***********001X",
            "1101011990*******",
            "1101011990*********",
            SYNTHETIC_FULL_ID,
        ):
            with self.subTest(invalid=invalid):
                data = copy.deepcopy(self.data)
                data["subjects"][0]["masked_id_number"] = invalid
                with self.assertRaises(network_workpaper.ValidationError):
                    network_workpaper.validate_run(
                        data,
                        workpaper_root=self.root,
                        formal_mode="draft",
                    )

        for key in ("cookie", "token", "password", "authorization"):
            with self.subTest(key=key):
                data = copy.deepcopy(self.data)
                data[key] = "not-a-real-value"
                with self.assertRaises(network_workpaper.ValidationError):
                    network_workpaper.validate_run(
                        data,
                        workpaper_root=self.root,
                        formal_mode="draft",
                    )

        with self.assertRaises(network_workpaper.ValidationError):
            network_workpaper.scan_sensitive_values({"note": SYNTHETIC_LEGACY_ID})

    def test_company_credit_code_is_field_aware_and_check_digit_validated(self) -> None:
        network_workpaper.validate_run(
            self.data,
            workpaper_root=self.root,
            formal_mode="final",
        )

        for invalid in (
            SYNTHETIC_VALID_CREDIT_CODE[:-1] + "1",
            "91110000TEST000001",
        ):
            with self.subTest(invalid=invalid):
                data = copy.deepcopy(self.data)
                data["subjects"][1]["credit_code"] = invalid
                with self.assertRaises(network_workpaper.ValidationError):
                    network_workpaper.validate_run(data, formal_mode="draft")

        natural_person = copy.deepcopy(self.data)
        natural_person["subjects"][0]["credit_code"] = SYNTHETIC_VALID_CREDIT_CODE
        with self.assertRaises(network_workpaper.ValidationError):
            network_workpaper.validate_run(natural_person, formal_mode="draft")

        for field in ("result_summary", "query_terms", "filters"):
            with self.subTest(field=field):
                data = copy.deepcopy(self.data)
                data["queries"][0][field] = SYNTHETIC_VALID_CREDIT_CODE
                with self.assertRaises(network_workpaper.ValidationError):
                    network_workpaper.validate_run(data, formal_mode="draft")

    def test_credit_code_prepare_validate_build_and_contextual_audit(self) -> None:
        subjects = copy.deepcopy(self.data["subjects"])
        subjects[1]["formal_identifier_mode"] = "auto_fill_company_credit_code"
        prepared = network_workpaper.prepare_run(
            {
                "run_id": "NQ-20260718-USCC",
                "timezone": "Asia/Shanghai",
                "profile": "custom",
                "project": copy.deepcopy(self.data["run"]["project"]),
                "formal_record": copy.deepcopy(self.data["run"]["formal_record"]),
                "subjects": subjects,
                "query_scope": {
                    "status": "user_confirmed",
                    "confirmed_at": "2026-07-18T10:00:00+08:00",
                    "selection_mode": "custom",
                    "items": [{
                        "scope_item_id": "SCOPE-001", "subject_id": "SUB-002",
                        "matter_category_id": "company.registration", "site_id": "registry",
                        "site_name": "匿名官方公示网站", "site_basis": "user_specified",
                        "allowed_domains": ["example.invalid"], "query_term_mode": "company_credit_code",
                        "conditions": [{"condition_id": "COND-001", "field": "period", "value": "2023年1月1日至查询日"}],
                    }],
                },
            }
        )
        network_workpaper.validate_run(prepared, formal_mode="draft")

        output_root = self.root / "credit-code-flow"
        paths = network_workpaper.build_outputs(
            prepared,
            workpaper_root=self.root,
            output_dir=output_root,
            template_docx=self.template,
            formal_mode="draft",
            layout="two-layer",
            overwrite=False,
        )
        docx_path = next(path for path in paths if path.suffix == ".docx")
        table_text = "\n".join(
            cell.text
            for table in Document(docx_path).tables
            for row in table.rows
            for cell in row.cells
        )
        self.assertIn(SYNTHETIC_VALID_CREDIT_CODE, table_text)

        contextual_report = network_workpaper.artifact_audit(
            docx_path,
            run_data=prepared,
        )
        self.assertTrue(contextual_report["ok"])
        self.assertEqual(contextual_report["validated_company_credit_codes"], 1)

        standalone_report = network_workpaper.artifact_audit(docx_path)
        self.assertFalse(standalone_report["ok"])
        self.assertEqual(standalone_report["validated_company_credit_codes"], 0)

    def test_prepare_rejects_prepopulated_queries(self) -> None:
        input_data = {
            "project": copy.deepcopy(self.data["run"]["project"]),
            "subjects": copy.deepcopy(self.data["subjects"]),
            "queries": [copy.deepcopy(self.data["queries"][0])],
        }
        with self.assertRaises(network_workpaper.ValidationError):
            network_workpaper.prepare_run(input_data)

    def test_natural_person_auto_fill_and_invalid_company_code_fail_closed(self) -> None:
        natural_person = copy.deepcopy(self.data)
        natural_person["subjects"][0]["formal_identifier_mode"] = "auto_fill_company_credit_code"
        with self.assertRaises(network_workpaper.ValidationError):
            network_workpaper.validate_run(natural_person, workpaper_root=self.root)

        invalid_company = copy.deepcopy(self.data)
        invalid_company["subjects"][1]["credit_code"] = SYNTHETIC_VALID_CREDIT_CODE[:-1] + "1"
        with self.assertRaises(network_workpaper.ValidationError):
            network_workpaper.validate_run(invalid_company, workpaper_root=self.root)

    def test_sensitive_candidate_in_query_or_manual_fill_artifact_fails_closed(self) -> None:
        for field in ("query_terms", "result_summary"):
            data = copy.deepcopy(self.data)
            data["queries"][0][field] = SYNTHETIC_VALID_CREDIT_CODE
            with self.subTest(field=field), self.assertRaises(network_workpaper.ValidationError):
                network_workpaper.validate_run(data, workpaper_root=self.root)

        data = copy.deepcopy(self.data)
        data["schema_version"] = "1.1"
        data["subjects"] = [copy.deepcopy(self.data["subjects"][1])]
        data["queries"] = []
        data["query_scope"] = {
            "status": "user_confirmed", "confirmed_at": "2026-07-19T10:00:00+08:00",
            "selection_mode": "custom", "items": [{
                "scope_item_id": "SCOPE-001", "subject_id": "SUB-002",
                "matter_category_id": "company.registration", "site_id": "registry",
                "site_name": "匿名官方公示网站", "site_basis": "user_specified",
                "allowed_domains": ["example.invalid"], "query_term_mode": "exact_subject_name",
                "conditions": [{"condition_id": "COND-001", "field": "period", "value": "2023年1月1日至查询日"}],
            }],
        }
        data["query_scope"]["items"][0]["conditions"][0]["value"] = SYNTHETIC_VALID_CREDIT_CODE
        with self.assertRaises(network_workpaper.ValidationError):
            network_workpaper.validate_run(data)

        archive = self.root / "manual-fill-archive"
        archive.mkdir()
        manual_docx = archive / "record.docx"
        make_template(manual_docx)
        document = Document(manual_docx)
        document.tables[0].cell(1, 2).text = SYNTHETIC_FULL_ID
        document.save(manual_docx)
        report = network_workpaper.artifact_audit(manual_docx)
        self.assertFalse(report["ok"])

    def test_v13_scope_preview_enforces_confirmed_domain_and_identifier_mode(self) -> None:
        data = copy.deepcopy(self.data)
        data["schema_version"] = "1.1"
        data["subjects"] = [copy.deepcopy(self.data["subjects"][1])]
        data["subjects"][0]["formal_identifier_mode"] = "auto_fill_company_credit_code"
        query = copy.deepcopy(self.data["queries"][0])
        query.update({
            "subject_id": "SUB-002", "site_id": "registry", "site_name": "匿名官方公示网站",
            "scope_item_id": "SCOPE-001", "matter_category_id": "company.registration",
            "query_term_mode": "company_credit_code", "condition_ids": ["COND-001"],
            "conditions": [{"condition_id": "COND-001", "field": "period", "value": "2023年1月1日至查询日"}],
            "url": "https://example.invalid/registry/search",
            "query_time": "2026-07-19T11:00:00+08:00",
        })
        query.pop("query_terms")
        query.pop("filters")
        data["queries"] = [query]
        data["query_scope"] = {
            "status": "user_confirmed", "confirmed_at": "2026-07-19T10:00:00+08:00",
            "selection_mode": "custom", "items": [{
                "scope_item_id": "SCOPE-001", "subject_id": "SUB-002",
                "matter_category_id": "company.registration", "site_id": "registry",
                "site_name": "匿名官方公示网站", "site_basis": "user_specified",
                "allowed_domains": ["example.invalid"], "query_term_mode": "company_credit_code",
                "conditions": [{"condition_id": "COND-001", "field": "period", "value": "2023年1月1日至查询日"}],
            }],
        }
        network_workpaper.validate_run(data, workpaper_root=self.root, formal_mode="final")
        self.assertEqual(
            network_workpaper.query_display_descriptions(
                "1.1", data["subjects"][0], data["queries"][0]
            ),
            ("经用户确认的统一社会信用代码", "period：2023年1月1日至查询日"),
        )
        for field in ("query_terms", "filters"):
            free_text = copy.deepcopy(data)
            free_text["queries"][0][field] = "未经确认、由 Agent 自行增加的条件"
            with self.subTest(field=field), self.assertRaises(network_workpaper.ValidationError):
                network_workpaper.validate_run(
                    free_text,
                    workpaper_root=self.root,
                    formal_mode="final",
                )
        before_confirmation = copy.deepcopy(data)
        before_confirmation["queries"][0]["query_time"] = "2026-07-19T09:59:59+08:00"
        with self.assertRaises(network_workpaper.ValidationError):
            network_workpaper.validate_run(
                before_confirmation,
                workpaper_root=self.root,
                formal_mode="final",
            )
        self.assertEqual(network_workpaper.scope_preview(data)["status"], "user_confirmed")
        self.assertEqual(network_workpaper.manual_identifier_fields_pending(data), [])
        paths = network_workpaper.build_outputs(
            data, workpaper_root=self.root, output_dir=self.root / "v13-auto",
            template_docx=self.template, formal_mode="final", layout="two-layer", overwrite=False,
        )
        internal_text = next(path for path in paths if path.suffix == ".md").read_text(encoding="utf-8")
        self.assertIn("查询条件：经用户确认的统一社会信用代码", internal_text)
        self.assertIn("筛选条件：period：2023年1月1日至查询日", internal_text)
        formal_table = Document(next(path for path in paths if path.suffix == ".docx")).tables[0]
        self.assertEqual(formal_table.rows[1].cells[2].text, SYNTHETIC_VALID_CREDIT_CODE)
        confirmed_query = data["queries"][0]
        data["queries"] = []
        data["query_scope"]["status"] = "pending_confirmation"
        data["query_scope"]["confirmed_at"] = ""
        self.assertEqual(network_workpaper.scope_preview(data)["status"], "pending_confirmation")
        data["queries"] = [confirmed_query]
        with self.assertRaises(network_workpaper.ValidationError):
            network_workpaper.validate_run(data, workpaper_root=self.root, formal_mode="draft")
        data["query_scope"]["status"] = "user_confirmed"
        data["query_scope"]["confirmed_at"] = "2026-07-19T10:00:00+08:00"
        data["queries"][0]["url"] = "https://outside.invalid/"
        with self.assertRaises(network_workpaper.ValidationError):
            network_workpaper.validate_run(data, workpaper_root=self.root, formal_mode="draft")

    def test_artifact_audit_allows_only_exact_context_code_and_never_filename(self) -> None:
        artifact_root = self.root / "credit-code-audit"
        artifact_root.mkdir()
        audit_run = copy.deepcopy(self.data)
        audit_run["subjects"][1]["formal_identifier_mode"] = "auto_fill_company_credit_code"
        (artifact_root / "record.md").write_text(
            SYNTHETIC_OTHER_VALID_CREDIT_CODE,
            encoding="utf-8",
        )
        report = network_workpaper.artifact_audit(
            artifact_root / "record.md", run_data=audit_run
        )
        self.assertFalse(report["ok"])

        (artifact_root / "record.md").write_text(
            SYNTHETIC_VALID_CREDIT_CODE,
            encoding="utf-8",
        )
        filename = artifact_root / f"{SYNTHETIC_VALID_CREDIT_CODE}.md"
        filename.write_text("匿名记录", encoding="utf-8")
        report = network_workpaper.artifact_audit(filename, run_data=audit_run)
        self.assertFalse(report["ok"])
        self.assertTrue(
            any("filename:" in issue for issue in report["issues"]),
            report["issues"],
        )

    def test_sensitive_page_may_be_recorded_without_screenshot(self) -> None:
        data = copy.deepcopy(self.data)
        query = data["queries"][0]
        query["screenshot_path"] = ""
        query["capture_kind"] = "not_retained_sensitive_page"
        query["result_summary"] = "页面显示匹配结果，因敏感信息保护未留存该页面截图"
        network_workpaper.validate_run(
            data,
            workpaper_root=self.root,
            formal_mode="draft",
        )

        query["result_summary"] = "页面显示匹配结果"
        query["follow_up"] = ""
        with self.assertRaises(network_workpaper.ValidationError):
            network_workpaper.validate_run(
                data,
                workpaper_root=self.root,
                formal_mode="draft",
            )

    def test_template_priority_and_markdown_only_fallback(self) -> None:
        explicit_root = self.root / "explicit"
        paths = network_workpaper.build_outputs(
            self.data,
            workpaper_root=self.root,
            output_dir=explicit_root,
            template_docx=self.template,
            formal_mode="draft",
            layout="two-layer",
            overwrite=False,
        )
        docx_path = next(path for path in paths if path.suffix == ".docx")
        explicit_text = "\n".join(p.text for p in Document(docx_path).paragraphs)
        self.assertIn("查询律师签署", explicit_text)
        self.assertNotIn("查询人员签名", explicit_text)

        original_bundled = network_workpaper.BUNDLED_TEMPLATE
        network_workpaper.BUNDLED_TEMPLATE = self.root / "missing-template.docx"
        try:
            markdown_root = self.root / "markdown-only"
            markdown_paths = network_workpaper.build_outputs(
                self.data,
                workpaper_root=self.root,
                output_dir=markdown_root,
                template_docx=None,
                formal_mode="draft",
                layout="two-layer",
                overwrite=False,
            )
            self.assertTrue(markdown_paths)
            self.assertTrue(all(path.suffix == ".md" for path in markdown_paths))

            with self.assertRaises(FileNotFoundError):
                network_workpaper.build_outputs(
                    self.data,
                    workpaper_root=self.root,
                    output_dir=self.root / "missing-final-template",
                    template_docx=None,
                    formal_mode="final",
                    layout="two-layer",
                    overwrite=False,
                )
        finally:
            network_workpaper.BUNDLED_TEMPLATE = original_bundled

    def test_bundled_template_is_generic_and_structurally_valid(self) -> None:
        report = network_workpaper.template_check(BUNDLED_TEMPLATE)
        self.assertTrue(report["privacy_safe"])
        with zipfile.ZipFile(BUNDLED_TEMPLATE) as archive:
            package_text = "\n".join(
                archive.read(name).decode("utf-8", errors="ignore")
                for name in archive.namelist()
                if name.endswith((".xml", ".rels"))
            )
        for forbidden in (
            "云亭",
            "律所地址",
            SYNTHETIC_FULL_ID,
            "/Users/",
            "OneDrive",
        ):
            self.assertNotIn(forbidden, package_text)
        self.assertRegex(package_text, r"\bPAGE\b")
        self.assertIn("统一社会信用代码/身份证号码", package_text)
        self.assertNotIn("身份证号码（脱敏）", package_text)

        document = Document(BUNDLED_TEMPLATE)
        properties = document.core_properties
        self.assertFalse(properties.author)
        self.assertFalse(properties.last_modified_by)
        self.assertFalse(properties.subject)
        self.assertFalse(properties.keywords)
        self.assertFalse(properties.comments)

    def test_doctor_reports_capability_degradation(self) -> None:
        self.assertEqual(
            network_workpaper.doctor_report("available")["mode"],
            "full-workflow",
        )
        self.assertEqual(
            network_workpaper.doctor_report("unavailable")["mode"],
            "local-processing-only",
        )

    def test_artifact_audit_checks_text_docx_and_png_metadata(self) -> None:
        clean_root = self.root / "clean-audit"
        clean_root.mkdir()
        (clean_root / "record.md").write_text("匿名核查记录", encoding="utf-8")
        make_template(clean_root / "record.docx")
        Image.new("RGB", (20, 20), "white").save(clean_root / "record.png")
        for clean_file in sorted(clean_root.iterdir()):
            with self.subTest(clean_file=clean_file.name):
                self.assertTrue(network_workpaper.artifact_audit(clean_file)["ok"])

        bad_root = self.root / "bad-audit"
        bad_root.mkdir()
        (bad_root / "record.md").write_text(
            "身份证号" + SYNTHETIC_FULL_ID, encoding="utf-8"
        )
        bad_docx = Document()
        bad_docx.add_paragraph("password=example")
        bad_docx.save(bad_root / "record.docx")
        metadata = PngImagePlugin.PngInfo()
        metadata.add_text("Subject", SYNTHETIC_FULL_ID)
        Image.new("RGB", (20, 20), "white").save(
            bad_root / "record.png", pnginfo=metadata
        )
        exif = Image.Exif()
        exif[0x010E] = SYNTHETIC_FULL_ID
        Image.new("RGB", (20, 20), "white").save(
            bad_root / "record.jpg", exif=exif
        )
        bad_reports = [
            network_workpaper.artifact_audit(path)
            for path in sorted(bad_root.iterdir())
        ]
        self.assertTrue(all(not report["ok"] for report in bad_reports))
        self.assertGreaterEqual(
            sum(len(report["issues"]) for report in bad_reports),
            4,
        )

    def test_template_and_audit_reject_sensitive_embedded_image_metadata(self) -> None:
        metadata = PngImagePlugin.PngInfo()
        metadata.add_text("Description", SYNTHETIC_FULL_ID)
        image_path = self.root / "embedded-sensitive.png"
        Image.new("RGB", (20, 20), "white").save(image_path, pnginfo=metadata)

        document = Document(self.template)
        document.add_paragraph().add_run().add_picture(str(image_path))
        template_path = self.root / "template-with-sensitive-image.docx"
        document.save(template_path)

        with self.assertRaises(network_workpaper.ValidationError):
            network_workpaper.template_check(template_path)
        self.assertFalse(network_workpaper.artifact_audit(template_path)["ok"])

    def test_subject_rows_replace_marker_in_place(self) -> None:
        template_path = self.root / "template-with-following-row.docx"
        make_template(template_path, trailing_subject_row=True)
        network_workpaper.template_check(template_path)
        output_path = self.root / "subject-row-output.docx"
        network_workpaper.build_formal_docx(
            self.data,
            template_path,
            output_path,
            "draft",
        )
        first_cells = [row.cells[0].text for row in Document(output_path).tables[0].rows]
        self.assertEqual(
            first_cells,
            ["主体角色", "自然人核查对象", "企业核查对象", "固定说明行"],
        )

    def test_overwrite_rejects_destination_symlinks(self) -> None:
        outside_target = self.root.parent / f"{self.root.name}-outside.txt"
        outside_target.write_text("ORIGINAL", encoding="utf-8")
        self.addCleanup(outside_target.unlink, missing_ok=True)

        output_root = self.root / "symlink-output"
        internal_dir = output_root / "01-内部底稿"
        internal_dir.mkdir(parents=True)
        build_link = internal_dir / "甲某网络核查底稿-20260717.md"
        try:
            build_link.symlink_to(outside_target)
        except OSError as exc:  # pragma: no cover - Windows without symlink permission
            self.skipTest(f"symbolic links are unavailable: {exc}")

        with self.assertRaises(network_workpaper.ValidationError):
            network_workpaper.build_outputs(
                self.data,
                workpaper_root=self.root,
                output_dir=output_root,
                template_docx=None,
                formal_mode="none",
                layout="two-layer",
                overwrite=True,
            )
        self.assertEqual(outside_target.read_text(encoding="utf-8"), "ORIGINAL")

        json_link = self.root / "prepared.json"
        json_link.symlink_to(outside_target)
        with self.assertRaises(network_workpaper.ValidationError):
            network_workpaper.write_json(json_link, {"safe": True}, overwrite=True)
        self.assertEqual(outside_target.read_text(encoding="utf-8"), "ORIGINAL")

        source = self.root / "symlink-source.png"
        Image.new("RGB", (20, 20), "white").save(source)
        watermark_link = self.root / "symlink-watermark.png"
        watermark_link.symlink_to(outside_target)
        with self.assertRaises(ValueError):
            watermark_capture.add_watermark(
                source,
                watermark_link,
                evidence_id="NQ-01-01-01",
                subject="甲某",
                source="示例网站",
                queried_at=watermark_capture.parse_timestamp("2026-07-17T09:10:11+08:00"),
                overwrite=True,
            )
        self.assertEqual(outside_target.read_text(encoding="utf-8"), "ORIGINAL")

    def test_watermark_preserves_timezone_and_rejects_sensitive_text(self) -> None:
        source = self.root / "source.png"
        output = self.root / "watermarked.png"
        source_image = Image.new("RGB", (640, 360), "white")
        source_draw = ImageDraw.Draw(source_image)
        source_draw.rectangle((20, 20, 620, 340), outline="black", width=4)
        source_draw.text((60, 150), "QUERY CONDITIONS / RESULTS", fill="black")
        source_image.save(source)
        queried_at = watermark_capture.parse_timestamp("2026-07-17T09:10:11+02:00")
        watermark_capture.add_watermark(
            source,
            output,
            evidence_id="NQ-01-01-01",
            subject="甲某",
            source="示例网站",
            queried_at=queried_at,
        )
        with Image.open(output) as image:
            self.assertEqual(image.info["QueriedAt"], "2026-07-17T15:10:11+08:00")
            self.assertEqual(image.info["CaptureKind"], "watermarked_page_only")

        with self.assertRaises(ValueError):
            watermark_capture.add_watermark(
                source,
                self.root / "bad.png",
                evidence_id="NQ-01-01-01",
                subject=SYNTHETIC_FULL_ID,
                source="示例网站",
                queried_at=queried_at,
            )
        with self.assertRaises(ValueError):
            watermark_capture.add_watermark(
                source,
                source,
                evidence_id="NQ-01-01-01",
                subject="甲某",
                source="示例网站",
                queried_at=queried_at,
                overwrite=True,
            )
        with self.assertRaises(ValueError):
            watermark_capture.add_watermark(
                source,
                self.root / "legacy-id.png",
                evidence_id="NQ-01-01-01",
                subject=SYNTHETIC_LEGACY_ID,
                source="示例网站",
                queried_at=queried_at,
            )

    def test_public_package_has_one_core_and_thin_adapter(self) -> None:
        plugin_root = PLUGIN_ROOT
        manifest = json.loads(
            (plugin_root / ".codex-plugin" / "plugin.json").read_text(encoding="utf-8")
        )
        self.assertEqual(manifest["version"], "2.0.1-beta")
        self.assertEqual(manifest["license"], "Apache-2.0")
        self.assertTrue((plugin_root / "LICENSE").is_file())
        self.assertTrue((plugin_root / "PRIVACY.md").is_file())
        self.assertEqual(
            list((plugin_root / "skills").glob("*/SKILL.md")),
            [SKILL_ROOT / "SKILL.md"],
        )

        skill_text = (SKILL_ROOT / "SKILL.md").read_text(encoding="utf-8")
        privacy_text = (plugin_root / "PRIVACY.md").read_text(encoding="utf-8")
        sensitive_rule_text = (
            SKILL_ROOT / "references" / "privacy-and-sensitive-query.md"
        ).read_text(encoding="utf-8")
        for policy_text in (privacy_text, sensitive_rule_text):
            self.assertIn("全部技术检查", policy_text)
            self.assertIn("自行补填", policy_text)
            self.assertIn("不得再次", policy_text)
        self.assertNotIn("Skill只生成脱敏版本", sensitive_rule_text)
        for forbidden in (
            "/Users/",
            "OneDrive",
            ".workbuddy/",
            ".codex/",
            "node_repl",
            "computer-use",
        ):
            self.assertNotIn(forbidden, skill_text)


if __name__ == "__main__":
    unittest.main()
