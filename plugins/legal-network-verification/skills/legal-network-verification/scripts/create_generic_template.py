#!/usr/bin/env python3
"""Build the neutral DOCX template bundled with the portable Skill."""

from __future__ import annotations

import argparse
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


CONTENT_WIDTH_DXA = 9030
TABLE_WIDTHS_DXA = (1800, 2700, 4530)
CJK_FONT = "Noto Serif CJK SC"


def set_run_font(run, size: float = 12, *, bold: bool | None = None) -> None:
    # Use one CJK-capable family for every script so that Word and LibreOffice
    # do not route Chinese text through a Latin-only font. Hosts without this
    # family can apply their normal document-font substitution.
    run.font.name = CJK_FONT
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), CJK_FONT)
    rfonts.set(qn("w:ascii"), CJK_FONT)
    rfonts.set(qn("w:hAnsi"), CJK_FONT)
    rfonts.set(qn("w:cs"), CJK_FONT)


def set_paragraph(paragraph, *, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line=False) -> None:
    paragraph.alignment = align
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(7.8)
    fmt.space_after = Pt(7.8)
    fmt.line_spacing = 1.5
    fmt.left_indent = None
    fmt.right_indent = None
    fmt.first_line_indent = Pt(24) if first_line else None


def set_cell_margins(cell) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", 80), ("bottom", 80), ("start", 120), ("end", 120)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "0")
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.find(qn(f"w:{edge}"))
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), "000000")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in TABLE_WIDTHS_DXA:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for cell, width in zip(row.cells, TABLE_WIDTHS_DXA):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def format_table_cell(cell, *, header: bool = False) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if header:
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "D9D9D9")
        cell._tc.get_or_add_tcPr().append(shading)
    for paragraph in cell.paragraphs:
        set_paragraph(paragraph, align=WD_ALIGN_PARAGRAPH.CENTER)
        paragraph.paragraph_format.space_before = Pt(2)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.line_spacing = 1
        paragraph.paragraph_format.first_line_indent = None
        for run in paragraph.runs:
            set_run_font(run, 10.5, bold=header)


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    set_run_font(run, 10.5)
    field_run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    field_run._r.extend((begin, instruction, separate, text, end))
    suffix = paragraph.add_run(" 页")
    set_run_font(suffix, 10.5)


def build_template(output: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.75)

    normal = document.styles["Normal"]
    normal.font.name = CJK_FONT
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    normal._element.rPr.rFonts.set(qn("w:ascii"), CJK_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), CJK_FONT)

    if "Network Query Label" not in document.styles:
        label_style = document.styles.add_style("Network Query Label", WD_STYLE_TYPE.CHARACTER)
        label_style.font.bold = True
        label_style.font.name = CJK_FONT
        label_style._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)

    title = document.add_paragraph()
    set_paragraph(title, align=WD_ALIGN_PARAGRAPH.CENTER)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(18)
    title_run = title.add_run("网络查询记录")
    set_run_font(title_run, 18, bold=True)

    metadata = (
        ("项目名称：", "【项目名称】"),
        ("查验事项：", "【查验事项】"),
        ("核查期间：", "【核查期间】"),
        ("查询时间：", "【查询时间】"),
        ("查询地点：", "【查询地点】"),
        ("查询人员：", "【查询人】"),
    )
    for label, value in metadata:
        paragraph = document.add_paragraph()
        set_paragraph(paragraph)
        label_run = paragraph.add_run(label)
        set_run_font(label_run, 12, bold=True)
        value_run = paragraph.add_run(value)
        set_run_font(value_run, 12)

    table = document.add_table(rows=2, cols=3)
    headers = ("主体角色", "查询对象", "身份标识（统一社会信用代码/身份证号码（脱敏））")
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
        format_table_cell(cell, header=True)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    table.rows[1].cells[0].text = "【SUBJECT_ROWS】"
    table.rows[1].cells[1].text = ""
    table.rows[1].cells[2].text = ""
    for cell in table.rows[1].cells:
        format_table_cell(cell)
    set_table_geometry(table)

    marker = document.add_paragraph("【QUERY_RESULT_ITEMS】")
    set_paragraph(marker, first_line=False)
    for run in marker.runs:
        set_run_font(run, 12)

    attachment = document.add_paragraph("附件：查询网址截图底稿另行保存。")
    set_paragraph(attachment)
    for run in attachment.runs:
        set_run_font(run, 12)

    signature = document.add_paragraph("查询人员签名：____________________")
    set_paragraph(signature, align=WD_ALIGN_PARAGRAPH.RIGHT)
    signature.paragraph_format.space_before = Pt(24)
    for run in signature.runs:
        set_run_font(run, 12)

    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    add_page_field(footer_paragraph)

    properties = document.core_properties
    properties.author = ""
    properties.last_modified_by = ""
    properties.title = "Generic Network Query Record Template"
    properties.subject = ""
    properties.keywords = ""
    properties.comments = ""
    properties.category = ""
    fixed_time = datetime(2000, 1, 1, tzinfo=timezone.utc)
    properties.created = fixed_time
    properties.modified = fixed_time

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output))


def main() -> int:
    parser = argparse.ArgumentParser(description="Build the generic network query DOCX template.")
    parser.add_argument("output", type=Path)
    parser.add_argument("--overwrite", action="store_true")
    args = parser.parse_args()
    if args.output.exists() and not args.overwrite:
        parser.error(f"output already exists: {args.output}")
    build_template(args.output)
    print(args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
