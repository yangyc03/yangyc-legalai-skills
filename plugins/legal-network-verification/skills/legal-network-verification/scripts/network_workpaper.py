#!/usr/bin/env python3
"""Prepare, validate, build and audit portable legal network workpapers.

This helper is deliberately offline. An authorized host Agent or the user performs browser interaction; this script
only handles local structured data, privacy checks and document generation.
"""

from __future__ import annotations

import argparse
import copy
import importlib.util
import json
import os
import platform
import re
import shutil
import sys
import tempfile
import zipfile
from collections import Counter
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any
from urllib.parse import parse_qsl, unquote, urlparse
from xml.etree import ElementTree as ET
from zoneinfo import ZoneInfo


SCHEMA_VERSION = "1.2"
COMPATIBLE_SCHEMA_VERSIONS = {"1.0", "1.1", "1.2"}
DEFAULT_TIMEZONE = "Asia/Shanghai"
ALLOWED_PROFILES = {"general-person", "general-company", "private-fund", "custom"}
ALLOWED_SUBJECT_TYPES = {"natural_person", "company"}
ALLOWED_STATUSES = {
    "identity_match",
    "no_match_displayed",
    "same_name_candidates",
    "not_applicable",
    "access_limited",
    "failed",
    "pending_review",
}
STATUS_LABELS = {
    "identity_match": "明确匹配",
    "no_match_displayed": "未显示可确认匹配",
    "same_name_candidates": "同名待核",
    "not_applicable": "查询对象不适用",
    "access_limited": "查询受限",
    "failed": "查询失败",
    "pending_review": "待复核",
}
SCREENSHOT_REQUIRED = {
    "identity_match",
    "no_match_displayed",
    "same_name_candidates",
    "not_applicable",
    "pending_review",
}
FORBIDDEN_KEYS = {
    "id_number",
    "identity_number",
    "identity_card_number",
    "birth_date",
    "birthday",
    "mobile",
    "phone",
    "email",
    "cookie",
    "token",
    "password",
    "authorization",
    "raw_response",
    "request_headers",
    "account",
    "account_name",
    "username",
    "profile_name",
    "session_id",
    "local_storage",
    "session_storage",
    "login_page_text",
    "raw_page_content",
}
FORBIDDEN_ABSOLUTE_PHRASES = {
    "确定不存在",
    "不存在任何",
    "无任何诉讼",
    "无任何执行",
    "无任何处罚",
    "无任何失信",
    "所有同名结果均与本人无关",
}
CHINA_ID_PATTERN = re.compile(
    r"(?<!\d)(?:\d{17}[0-9Xx]|\d{8}(?:0[1-9]|1[0-2])(?:0[1-9]|[12]\d|3[01])\d{3})(?!\d)"
)
USCC_ALPHABET = "0123456789ABCDEFGHJKLMNPQRTUWXY"
USCC_PATTERN = re.compile(r"^[0-9ABCDEFGHJKLMNPQRTUWXY]{18}$")
USCC_TOKEN_PATTERN = re.compile(
    r"(?<![0-9A-Z])[0-9ABCDEFGHJKLMNPQRTUWXY]{18}(?![0-9A-Z])",
    re.IGNORECASE,
)
USCC_WEIGHTS = (1, 3, 9, 27, 19, 26, 16, 17, 20, 29, 25, 13, 8, 24, 10, 30, 28)
BIRTHDATE_LABEL_PATTERN = re.compile(
    r"(?:出生日期|出生年月|生日)\s*[:：]?\s*(?:19|20)\d{2}[-/.年]\d{1,2}(?:[-/.月]\d{1,2}日?)?"
)
ZERO_WIDTH_SPACE = "\u200b"
MASKED_ID_PATTERN = re.compile(r"^\d{10}\*{8}$")
SENSITIVE_NO_CAPTURE_KIND = "not_retained_sensitive_page"
SENSITIVE_NO_CAPTURE_EXPLANATION = "因敏感信息保护未留存该页面截图"
WATERMARKED_CAPTURE_KIND = "watermarked_page_only"
OUTPUT_DOCX_FONT = "Noto Serif CJK SC"
BUNDLED_TEMPLATE = (
    Path(__file__).resolve().parents[1] / "assets" / "generic-network-query-record.docx"
)
TEXT_ARTIFACT_SUFFIXES = {
    ".csv", ".json", ".log", ".md", ".rels", ".rst", ".svg", ".tsv", ".txt", ".xml", ".yaml", ".yml"
}
RASTER_IMAGE_SUFFIXES = {".bmp", ".gif", ".jpeg", ".jpg", ".png", ".tif", ".tiff", ".webp"}
LOCAL_PATH_PATTERNS = (
    re.compile(r"/Users/[^/\s]+/"),
    re.compile(r"[A-Za-z]:\\Users\\[^\\\s]+\\"),
    re.compile(r"(?:^|[/\\])\.codex(?:[/\\]|$)"),
    re.compile(r"(?:^|[/\\])\.workbuddy(?:[/\\]|$)"),
    re.compile(r"OneDrive", re.IGNORECASE),
)
SENSITIVE_LABEL_PATTERN = re.compile(
    r"(?:cookie|password|token|authorization|raw_response|request_headers)[\"']?\s*[:=]",
    re.IGNORECASE,
)
ABSOLUTE_LOCAL_PATH_PATTERN = re.compile(
    r"(?<![A-Za-z0-9])(?:/(?:Users|private|tmp|Volumes)/|[A-Za-z]:\\)",
    re.IGNORECASE,
)
SENSITIVE_URL_NAME_PATTERN = re.compile(
    r"(?:^|[_\-.])(?:token|session(?:_id)?|sid|auth(?:orization)?|ticket|cookie|password|credential|api[_-]?key)(?:$|[_\-.])",
    re.IGNORECASE,
)
SENSITIVE_URL_PATH_PATTERN = re.compile(
    r"(?:^|/)(?:access[_-]?token|refresh[_-]?token|token|j?session(?:[_-]?id)?|sid|auth(?:orization|entication|[_-]?token)?|code|ticket|api[_-]?key|key|cookie|password|credential)(?:/|$)",
    re.IGNORECASE,
)
FORMAL_IDENTIFIER_MODES = {"user_fill", "auto_fill_company_credit_code"}
QUERY_SCOPE_STATUS = "user_confirmed"
ALLOWED_EXECUTION_STATES = {
    "pending_user_action",
    "ready",
    "running",
    "completed",
    "blocked",
    "failed",
}
TERMINAL_EXECUTION_STATES = {"completed", "blocked", "failed"}
ALLOWED_ACCESS_STATES = {
    "unknown",
    "not_required",
    "authenticated",
    "login_required",
    "session_expired",
    "human_verification_required",
    "limited",
}
ALLOWED_SUBMIT_STATES = {"not_submitted", "submitted", "unknown"}
ALLOWED_BROWSER_SURFACES = {"in_app_browser", "chrome", "computer_use"}
ALLOWED_BROWSER_SELECTIONS = {
    "user_explicit",
    "automatic_primary",
    "automatic_fallback",
}
ALLOWED_BLOCK_REASONS = {
    "login_required",
    "session_expired",
    "human_verification_required",
    "rate_limited",
    "permission_required",
    "paywall",
    "site_maintenance",
    "network_error",
    "tool_error",
    "submit_state_unknown",
    "user_stopped",
}
ALLOWED_RESULT_SIGNALS = {
    "not_observed",
    "explicit_zero_count",
    "explicit_empty_state",
    "matching_results_displayed",
    "same_name_candidates_displayed",
    "not_applicable",
}
ALLOWED_CAPTURE_REVIEWS = {
    "not_reviewed",
    "conditions_and_result_visible",
    "sensitive_no_capture",
}
ALLOWED_ATTEMPT_EVENTS = {
    "access_probe",
    "browser_fallback",
    "handoff_started",
    "user_resume_confirmed",
    "resume_safety_check_passed",
    "page_read",
    "click",
    "screenshot",
    "query_submitted",
    "result_region_loaded",
    "login_required",
    "session_expired",
    "human_verification_required",
    "rate_limited",
    "permission_required",
    "paywall",
    "site_maintenance",
    "network_error",
    "tool_error",
    "explicit_resubmit_confirmed",
    "sensitive_result_review_confirmed",
    "user_stopped",
}
HANDOFF_START_EVENTS = {
    "login_required",
    "session_expired",
    "human_verification_required",
    "handoff_started",
}
HANDOFF_FORBIDDEN_EVENTS = {
    "access_probe",
    "browser_fallback",
    "page_read",
    "click",
    "screenshot",
    "query_submitted",
    "result_region_loaded",
}
POST_SUBMIT_CHALLENGE_EVENTS = {
    "login_required",
    "session_expired",
    "human_verification_required",
}
POST_SUBMIT_UNCERTAIN_EVENTS = POST_SUBMIT_CHALLENGE_EVENTS | {"network_error", "tool_error"}
BLOCKED_REASONS = {
    "rate_limited",
    "permission_required",
    "paywall",
    "site_maintenance",
    "user_stopped",
}
FAILED_REASONS = {"network_error", "tool_error"}
RESTRICTED_NEGATIVE_SIGNALS = {
    "未发现记录",
    "未显示匹配",
    "没有记录",
    "零条",
    "0条",
    "0 条",
    "暂无符合",
    "无匹配",
    "空结果",
}
CANONICAL_RESTRICTED_RESULT = "本项因网站访问条件受限未完整完成，暂不作否定性结论。"
CANONICAL_FAILED_RESULT = "本项因技术原因未完成，暂不作否定性结论。"
CANONICAL_RESTRICTED_ASSESSMENT = "本项不支持作出否定性结论。"
CANONICAL_RESTRICTED_FOLLOW_UP = "建议另行补充核验。"
EXECUTION_STATUS_LABELS = {
    "not_started": "未开始",
    "pending_login": "待登录",
    "pending_verification": "待人工验证",
    "submit_unknown": "提交状态待确认",
    "in_progress": "核查进行中",
    "completed_match": "已完成有结果",
    "completed_no_match": "已完成零结果",
    "same_name": "同名待排除",
    "access_limited": "访问受限",
    "failed": "技术失败",
    "not_applicable": "不适用",
    "pending_review": "已完成有结果（待复核）",
}
COMPANY_DEFAULT_CATEGORIES = (
    "诉讼及裁判文书",
    "法院公告",
    "执行",
    "失信",
    "限高",
    "破产及相关程序",
    "行政处罚",
    "经营异常及严重违法失信",
    "司法协助及股权冻结",
    "许可证或资质异常",
)
COMPANY_CONFIRM_CATEGORIES = ("监管措施/市场禁入/行业自律处分", "重大税收违法/欠税公告")
NATURAL_PERSON_DEFAULT_CATEGORIES = ("裁判文书", "法院公告", "执行", "失信", "限高")
NATURAL_PERSON_CONFIRM_CATEGORIES = ("行政/行业监管", "市场禁入/自律处分")
DEDICATED_ARTIFACT_DIR_NAMES = {"01-内部底稿", "02-正式记录"}
INTERNAL_ARTIFACT_DIR_NAME = "01-内部底稿"
FORMAL_ARTIFACT_DIR_NAME = "02-正式记录"
RUN_FILE_NAME = "network-verification.json"
SCREENSHOT_DIR_NAME = "截图"
SUPPORTED_EXPLICIT_ARTIFACT_SUFFIXES = {".json", ".md", ".docx"} | RASTER_IMAGE_SUFFIXES
WORDPROCESSINGML_NAMESPACE = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
WORD_TEXT_TAGS = {
    f"{{{WORDPROCESSINGML_NAMESPACE}}}t",
    f"{{{WORDPROCESSINGML_NAMESPACE}}}delText",
    f"{{{WORDPROCESSINGML_NAMESPACE}}}instrText",
}
DOCX_WORD_TEXT_PART_PATTERN = re.compile(
    r"^word/(?:document|footnotes|endnotes|comments)\.xml$|"
    r"^word/(?:header|footer)\d+\.xml$|^word/glossary/document\.xml$"
)


class ValidationError(ValueError):
    """Raised when a run file violates the Skill contract."""


class ArtifactAuditPlan:
    """A closed audit set derived from one validated network-verification run."""

    def __init__(
        self,
        *,
        files: list[Path],
        audit_root: Path | None,
        pre_audited_files: set[Path],
        preflight_issues: list[str],
        unexpected_artifacts: list[str],
    ) -> None:
        self.files = files
        self.audit_root = audit_root
        self.pre_audited_files = pre_audited_files
        self.preflight_issues = preflight_issues
        self.unexpected_artifacts = unexpected_artifacts


def load_json(path: Path) -> dict[str, Any]:
    if not path.is_file():
        raise FileNotFoundError(f"JSON file not found: {path}")
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise ValidationError("JSON root must be an object.")
    return value


def write_json(path: Path, value: dict[str, Any], *, overwrite: bool = False) -> None:
    if path.is_symlink():
        raise ValidationError(f"Output path must not be a symbolic link: {path}")
    if path.exists() and not overwrite:
        raise FileExistsError(f"Output file already exists: {path}")
    path.parent.mkdir(parents=True, exist_ok=True)
    atomic_write_text(path, json.dumps(value, ensure_ascii=False, indent=2) + "\n")


def atomic_write_text(path: Path, text: str) -> None:
    file_descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{path.name}.",
        suffix=".tmp",
        dir=path.parent,
    )
    temporary_path = Path(temporary_name)
    try:
        with os.fdopen(file_descriptor, "w", encoding="utf-8") as handle:
            handle.write(text)
        os.replace(temporary_path, path)
    finally:
        temporary_path.unlink(missing_ok=True)


def safe_output_destination(path: Path, root: Path, label: str) -> Path:
    if path.is_symlink():
        raise ValidationError(f"{label} must not be a symbolic link: {path}")
    resolved_parent = path.parent.resolve()
    try:
        resolved_parent.relative_to(root)
    except ValueError as exc:
        raise ValidationError(f"{label} escapes workpaper_root: {path}") from exc
    destination = resolved_parent / path.name
    if destination.is_symlink():
        raise ValidationError(f"{label} must not be a symbolic link: {path}")
    return destination


def atomic_copy_file(source: Path, destination: Path) -> None:
    file_descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{destination.name}.",
        suffix=".tmp",
        dir=destination.parent,
    )
    os.close(file_descriptor)
    temporary_path = Path(temporary_name)
    try:
        shutil.copyfile(source, temporary_path)
        os.replace(temporary_path, destination)
    finally:
        temporary_path.unlink(missing_ok=True)


def require_object(value: Any, label: str) -> dict[str, Any]:
    if not isinstance(value, dict):
        raise ValidationError(f"{label} must be an object.")
    return value


def require_list(value: Any, label: str) -> list[Any]:
    if not isinstance(value, list):
        raise ValidationError(f"{label} must be a list.")
    return value


def require_text(value: Any, label: str) -> str:
    text = str(value or "").strip()
    if not text:
        raise ValidationError(f"{label} is required.")
    return text


def require_bool(value: Any, label: str) -> bool:
    if not isinstance(value, bool):
        raise ValidationError(f"{label} must be a boolean.")
    return value


def reject_unknown_keys(value: dict[str, Any], allowed: set[str], label: str) -> None:
    unknown = sorted(set(value) - allowed)
    if unknown:
        raise ValidationError(f"{label} contains unknown field(s): {', '.join(unknown)}")


def require_timestamp(value: Any, label: str) -> datetime:
    text = require_text(value, label)
    try:
        parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
    except ValueError as exc:
        raise ValidationError(f"{label} must be an ISO 8601 timestamp.") from exc
    if parsed.tzinfo is None:
        raise ValidationError(f"{label} must include a timezone offset.")
    return parsed


def validate_company_credit_code(value: Any, label: str) -> str:
    code = require_text(value, label)
    if not USCC_PATTERN.fullmatch(code):
        raise ValidationError(
            f"{label} must be an 18-character unified social credit code using the GB 32100 character set."
        )
    total = sum(
        USCC_ALPHABET.index(character) * weight
        for character, weight in zip(code[:17], USCC_WEIGHTS)
    )
    expected = USCC_ALPHABET[(31 - total % 31) % 31]
    if code[-1] != expected:
        raise ValidationError(f"{label} has an invalid unified social credit code check character.")
    return code


def valid_company_credit_code_tokens(text: str) -> list[str]:
    tokens: list[str] = []
    for match in USCC_TOKEN_PATTERN.finditer(text.upper()):
        token = match.group(0)
        try:
            validate_company_credit_code(token, "company_credit_code")
        except ValidationError:
            continue
        tokens.append(token)
    return tokens


def company_credit_code_locations(value: Any, location: str) -> dict[str, str]:
    if not isinstance(value, dict):
        return {}
    subjects = value.get("subjects")
    if not isinstance(subjects, list):
        return {}
    allowed: dict[str, str] = {}
    for index, raw_subject in enumerate(subjects):
        if not isinstance(raw_subject, dict) or raw_subject.get("type") != "company":
            continue
        credit_code = str(raw_subject.get("credit_code") or "").strip()
        if not credit_code:
            continue
        credit_code_location = f"{location}.subjects[{index}].credit_code"
        allowed[credit_code_location] = validate_company_credit_code(
            credit_code,
            credit_code_location,
        )
    return allowed


def scan_sensitive_values(
    value: Any,
    location: str = "run",
    *,
    allowed_company_credit_code_locations: dict[str, str] | None = None,
) -> None:
    if allowed_company_credit_code_locations is None:
        allowed_company_credit_code_locations = company_credit_code_locations(value, location)
    if isinstance(value, dict):
        for key, item in value.items():
            normalized = str(key).strip().lower().replace("-", "_")
            if normalized in FORBIDDEN_KEYS:
                raise ValidationError(f"Sensitive field is not allowed: {location}.{key}")
            scan_sensitive_values(
                item,
                f"{location}.{key}",
                allowed_company_credit_code_locations=allowed_company_credit_code_locations,
            )
        return
    if isinstance(value, list):
        for index, item in enumerate(value):
            scan_sensitive_values(
                item,
                f"{location}[{index}]",
                allowed_company_credit_code_locations=allowed_company_credit_code_locations,
            )
        return
    if isinstance(value, str):
        allowed_credit_code = allowed_company_credit_code_locations.get(location)
        if allowed_credit_code is not None and value.strip() == allowed_credit_code:
            return
        if CHINA_ID_PATTERN.search(value):
            raise ValidationError(f"Full identity number is not allowed: {location}")
        if valid_company_credit_code_tokens(value):
            raise ValidationError(f"Company credit code is not allowed at this location: {location}")
        if BIRTHDATE_LABEL_PATTERN.search(value):
            raise ValidationError(f"Exact birth date must not be persisted: {location}")


def validate_persisted_text_safety(value: Any, location: str = "run") -> None:
    """Reject machine-local paths and credential-shaped text in schema 1.2 data."""
    if isinstance(value, dict):
        for key, item in value.items():
            validate_persisted_text_safety(item, f"{location}.{key}")
        return
    if isinstance(value, list):
        for index, item in enumerate(value):
            validate_persisted_text_safety(item, f"{location}[{index}]")
        return
    if not isinstance(value, str):
        return
    text = value.strip()
    if ABSOLUTE_LOCAL_PATH_PATTERN.search(text):
        raise ValidationError(f"Machine-local absolute path is not allowed: {location}")
    if SENSITIVE_LABEL_PATTERN.search(text):
        raise ValidationError(f"Credential-shaped text is not allowed: {location}")


def validate_safe_relative_path(value: Any, label: str) -> Path:
    text = require_text(value, label)
    path = Path(text)
    if path.is_absolute() or ".." in path.parts:
        raise ValidationError(f"{label} must be a safe relative path without '..'.")
    return path


def validate_archival_url(
    value: Any,
    label: str,
    *,
    allowed_domains: list[str],
    allow_query: bool,
) -> str:
    text = require_text(value, label)
    if any(character.isspace() or ord(character) < 32 for character in text):
        raise ValidationError(f"{label} contains whitespace or control characters.")
    try:
        parsed = urlparse(text)
        parsed.port
    except ValueError as exc:
        raise ValidationError(f"{label} has an invalid network structure.") from exc
    if parsed.scheme.lower() not in {"http", "https"} or not parsed.hostname:
        raise ValidationError(f"{label} must be an absolute HTTP(S) URL.")
    if parsed.username is not None or parsed.password is not None or "@" in parsed.netloc:
        raise ValidationError(f"{label} must not contain username or password.")
    if not hostname_allowed(parsed.hostname, allowed_domains):
        raise ValidationError(f"{label} domain exceeds confirmed scope.")
    if parsed.fragment:
        raise ValidationError(f"{label} must not retain a fragment.")
    if parsed.params:
        raise ValidationError(f"{label} must not retain semicolon parameters.")
    if not allow_query and parsed.query:
        raise ValidationError(f"{label} must not contain a query string.")
    normalized_path = parsed.path
    for _ in range(3):
        decoded = unquote(normalized_path)
        if decoded == normalized_path:
            break
        normalized_path = decoded
    if "%" in normalized_path:
        raise ValidationError(f"{label} contains unresolved percent encoding.")
    if any(character.isspace() or ord(character) < 32 for character in normalized_path):
        raise ValidationError(f"{label} contains encoded whitespace or control characters.")
    if "\\" in normalized_path:
        raise ValidationError(f"{label} must not contain backslash path separators.")
    if any(delimiter in normalized_path for delimiter in (";", "=", ":")):
        raise ValidationError(f"{label} must not retain path parameters or assignments.")
    if SENSITIVE_URL_PATH_PATTERN.search(normalized_path):
        raise ValidationError(f"{label} contains a credential-shaped path segment.")
    for name, _ in parse_qsl(parsed.query, keep_blank_values=True):
        if SENSITIVE_URL_NAME_PATTERN.search(name):
            raise ValidationError(f"{label} contains a sensitive query parameter name.")
    return text


def check_wording(text: str, label: str) -> None:
    for phrase in FORBIDDEN_ABSOLUTE_PHRASES:
        if phrase in text:
            raise ValidationError(f"Absolute nonexistence wording is not allowed in {label}: {phrase}")


def text_privacy_issues(text: str, label: str) -> list[str]:
    issues: list[str] = []
    if CHINA_ID_PATTERN.search(text):
        issues.append(f"full identity number: {label}")
    if valid_company_credit_code_tokens(text):
        issues.append(f"company credit code outside approved context: {label}")
    if SENSITIVE_LABEL_PATTERN.search(text):
        issues.append(f"credential or authorization field: {label}")
    for pattern in LOCAL_PATH_PATTERNS:
        if pattern.search(text):
            issues.append(f"machine-specific path: {label}")
            break
    return issues


def text_metadata_privacy_issues(text: str, label: str) -> list[str]:
    """Check non-identity privacy markers without granting any code exception."""
    issues: list[str] = []
    if SENSITIVE_LABEL_PATTERN.search(text):
        issues.append(f"credential or authorization field: {label}")
    for pattern in LOCAL_PATH_PATTERNS:
        if pattern.search(text):
            issues.append(f"machine-specific path: {label}")
            break
    return issues


def replace_approved_docx_company_codes(
    payload: bytes,
    run_data: dict[str, Any] | None,
) -> bytes:
    """Mask codes only in the template's subject identity column before audit."""
    if run_data is None:
        return payload
    approved_rows = Counter(
        (
            str(subject.get("role") or "").strip(),
            str(subject.get("name") or "").strip(),
            str(subject.get("credit_code") or "").strip(),
        )
        for subject in require_list(run_data.get("subjects"), "subjects")
        if subject.get("type") == "company"
        and str(subject.get("formal_identifier_mode") or "user_fill").strip()
        == "auto_fill_company_credit_code"
        and str(subject.get("credit_code") or "").strip()
    )
    if not approved_rows:
        return payload
    try:
        root = ET.fromstring(payload)
    except ET.ParseError:
        return payload
    row_tag = f"{{{WORDPROCESSINGML_NAMESPACE}}}tr"
    cell_tag = f"{{{WORDPROCESSINGML_NAMESPACE}}}tc"
    text_tags = WORD_TEXT_TAGS
    matched_rows: Counter[tuple[str, str, str]] = Counter()
    row_nodes: dict[tuple[str, str, str], list[list[Any]]] = {}
    for table in root.iter(f"{{{WORDPROCESSINGML_NAMESPACE}}}tbl"):
        rows = list(table.iter(row_tag))
        header_indexes: set[int] = set()
        for row_index, row in enumerate(rows):
            cells = list(row.iter(cell_tag))
            cell_texts = [
                "".join(node.text or "" for node in cell.iter() if node.tag in text_tags)
                for cell in cells
            ]
            if (
                len(cell_texts) >= 3
                and "主体角色" in cell_texts[0]
                and any(label in cell_texts[1] for label in ("姓名/名称", "查询对象"))
                and any(
                    label in cell_texts[2]
                    for label in ("统一社会信用代码/身份证号码", "身份号码")
                )
            ):
                header_indexes.add(row_index)
        for header_index in header_indexes:
            for row in rows[header_index + 1 :]:
                cells = list(row.iter(cell_tag))
                if len(cells) < 3:
                    continue
                cell = cells[2]
                nodes = [node for node in cell.iter() if node.tag in text_tags]
                cell_text = "".join(node.text or "" for node in nodes)
                row_key = (
                    "".join(node.text or "" for node in cells[0].iter() if node.tag in text_tags),
                    "".join(node.text or "" for node in cells[1].iter() if node.tag in text_tags),
                    cell_text,
                )
                if row_key not in approved_rows:
                    continue
                matched_rows[row_key] += 1
                row_nodes.setdefault(row_key, []).append(nodes)
    if any(matched_rows[key] != expected_count for key, expected_count in approved_rows.items()):
        return payload
    for row_key, node_groups in row_nodes.items():
        for nodes in node_groups:
            if nodes:
                nodes[0].text = "[VERIFIED_COMPANY_CODE]"
                for node in nodes[1:]:
                    node.text = ""
    return ET.tostring(root, encoding="utf-8")


def metadata_value_text(value: Any) -> str:
    if isinstance(value, bytes):
        decoded: list[str] = []
        for encoding in ("utf-8", "utf-16-le", "utf-16-be", "latin-1"):
            try:
                text = value.decode(encoding)
            except UnicodeDecodeError:
                continue
            if text not in decoded:
                decoded.append(text)
        return "\n".join(decoded)
    if isinstance(value, dict):
        return "\n".join(
            f"{metadata_value_text(key)}={metadata_value_text(item)}" for key, item in value.items()
        )
    if isinstance(value, (list, tuple, set)):
        return "\n".join(metadata_value_text(item) for item in value)
    return str(value)


def image_metadata_privacy_issues(
    payload: bytes,
    label: str,
) -> list[str]:
    try:
        from PIL import Image

        with Image.open(BytesIO(payload)) as image:
            metadata_items: list[tuple[Any, Any]] = list(image.info.items())
            try:
                metadata_items.extend(image.getexif().items())
            except (AttributeError, OSError, ValueError):
                pass
            metadata_text = "\n".join(
                f"{metadata_value_text(key)}={metadata_value_text(value)}"
                for key, value in metadata_items
            )
    except Exception as exc:
        return [f"unreadable image metadata: {label}: {exc}"]
    return text_privacy_issues(metadata_text, label)


def docx_text_container_privacy_issues(
    payload: bytes,
    label: str,
) -> list[str]:
    """Join Word text runs inside each paragraph before checking persisted text."""
    try:
        root = ET.fromstring(payload)
    except ET.ParseError as exc:
        return [f"unreadable Word text part: {label}: {exc}"]
    issues: list[str] = []
    paragraph_tag = f"{{{WORDPROCESSINGML_NAMESPACE}}}p"
    for index, paragraph in enumerate(root.iter(paragraph_tag), start=1):
        text = "".join(
            item.text or "" for item in paragraph.iter() if item.tag in WORD_TEXT_TAGS
        )
        issues.extend(
            text_privacy_issues(text, f"{label}:paragraph[{index}]")
        )
    return issues


def audit_date_token(data: dict[str, Any]) -> str:
    queries = require_list(data.get("queries"), "queries")
    if not queries:
        return "未查询"
    latest = max(require_timestamp(item["query_time"], "query.query_time") for item in queries)
    return latest.astimezone(ZoneInfo(data["run"]["timezone"])).strftime("%Y%m%d")


def expected_internal_artifact_paths(data: dict[str, Any]) -> set[Path]:
    """Derive the only files a two-layer internal directory may contain."""
    expected = {Path(INTERNAL_ARTIFACT_DIR_NAME) / RUN_FILE_NAME}
    screenshot_prefix = Path(INTERNAL_ARTIFACT_DIR_NAME) / SCREENSHOT_DIR_NAME
    for index, query in enumerate(require_list(data.get("queries"), "queries")):
        screenshot_text = str(query.get("screenshot_path") or "").strip()
        if not screenshot_text:
            continue
        screenshot_path = validate_safe_relative_path(
            screenshot_text, f"queries[{index}].screenshot_path"
        )
        try:
            screenshot_path.relative_to(screenshot_prefix)
        except ValueError as exc:
            raise ValidationError(
                "Two-layer artifact-audit accepts screenshots only under "
                "01-内部底稿/截图/."
            ) from exc
        if screenshot_path.suffix.lower() not in RASTER_IMAGE_SUFFIXES:
            raise ValidationError(
                "Two-layer artifact-audit accepts only raster watermarked screenshots: "
                + screenshot_path.as_posix()
            )
        expected.add(screenshot_path)

    subjects = require_list(data.get("subjects"), "subjects")
    name_counts: dict[str, int] = {}
    for subject in subjects:
        name = subject["name"]
        name_counts[name] = name_counts.get(name, 0) + 1
    date_token = audit_date_token(data)
    for subject in subjects:
        suffix = f"-{subject['subject_id']}" if name_counts[subject["name"]] > 1 else ""
        expected.add(
            Path(INTERNAL_ARTIFACT_DIR_NAME)
            / f"{safe_filename(subject['name'])}网络核查底稿-{date_token}{suffix}.md"
        )
    return expected


def expected_formal_artifact_paths(data: dict[str, Any]) -> set[Path]:
    """Derive the draft/final formal names that this script can generate."""
    project = require_object(data.get("run"), "run")["project"]
    project_name = safe_filename(project["short_name"])
    date_token = audit_date_token(data)
    return {
        Path(FORMAL_ARTIFACT_DIR_NAME) / f"网络查询记录-{project_name}-{date_token}.docx",
        Path(FORMAL_ARTIFACT_DIR_NAME) / f"网络查询记录-{project_name}-{date_token}_草稿.docx",
    }


def ensure_regular_project_file(path: Path, project_root: Path, label: str) -> Path:
    """Reject symlinks and any candidate that resolves outside this project root."""
    try:
        relative = path.relative_to(project_root)
    except ValueError as exc:
        raise ValidationError(f"Artifact is outside the project root: {label}") from exc
    current = project_root
    for part in relative.parts:
        current = current / part
        if current.is_symlink():
            raise ValidationError(f"Symlinked artifact paths are not allowed: {label}")
    if not path.is_file():
        raise FileNotFoundError(f"Artifact file not found: {label}")
    resolved = path.resolve()
    try:
        resolved.relative_to(project_root)
    except ValueError as exc:
        raise ValidationError(f"Artifact escapes the project root: {label}") from exc
    return resolved


def dedicated_project_root(target: Path) -> tuple[Path, Path]:
    """Return canonical project/layer paths without accepting a symlinked layer."""
    if not target.exists():
        raise FileNotFoundError(f"Artifact path not found: {target}")
    if target.name not in DEDICATED_ARTIFACT_DIR_NAMES:
        raise ValidationError(
            "artifact-audit accepts only a dedicated 01-内部底稿/02-正式记录 directory "
            "or one explicit artifact file; do not scan a client project root or broad directory."
        )
    if not target.is_dir() or target.is_symlink():
        raise ValidationError("artifact-audit requires a real dedicated generated directory, not a symlink.")
    project_root = target.parent.resolve()
    layer = target.resolve()
    if layer.parent != project_root:
        raise ValidationError("Dedicated generated directory must be directly inside its project root.")
    return project_root, layer


def load_audit_run(run_path: Path, project_root: Path) -> tuple[dict[str, Any], list[str]]:
    """Read and validate the one run file that authorizes a directory audit."""
    label = run_path.relative_to(project_root).as_posix()
    ensure_regular_project_file(run_path, project_root, label)
    text = run_path.read_text(encoding="utf-8")
    try:
        data = json.loads(text)
    except json.JSONDecodeError as exc:
        raise ValidationError(f"Invalid network-verification.json: {label}") from exc
    if not isinstance(data, dict):
        raise ValidationError(f"network-verification.json root must be an object: {label}")
    validate_run(data, workpaper_root=project_root, formal_mode="draft")
    # The run file is checked structurally: only subjects[].credit_code may
    # contain a validated company code. Other serialized locations remain
    # blocked even when the same value is valid for a company subject.
    scan_sensitive_values(data)
    return data, text_metadata_privacy_issues(text, label)


def enumerate_layer_files(layer: Path) -> set[str]:
    """List relative file names only; never open unapproved artifact content."""
    relative_names: set[str] = set()
    for current_text, dirnames, filenames in os.walk(layer, followlinks=False):
        current = Path(current_text)
        for dirname in dirnames:
            candidate = current / dirname
            if candidate.is_symlink():
                relative = candidate.relative_to(layer).as_posix()
                raise ValidationError(f"Symlinked artifact paths are not allowed: {relative}")
        for filename in filenames:
            relative_names.add((current / filename).relative_to(layer).as_posix())
    return relative_names


def directory_audit_plan(target: Path) -> ArtifactAuditPlan:
    """Build a closed file plan from the run file before examining any other file."""
    project_root, layer = dedicated_project_root(target)
    internal_layer = project_root / INTERNAL_ARTIFACT_DIR_NAME
    run_path = internal_layer / RUN_FILE_NAME
    if not internal_layer.is_dir() or internal_layer.is_symlink():
        raise ValidationError(
            "artifact-audit requires sibling 01-内部底稿/network-verification.json for this project."
        )
    if not run_path.is_file() or run_path.is_symlink():
        raise ValidationError(
            "artifact-audit requires a direct 01-内部底稿/network-verification.json run file."
        )

    data, preflight_issues = load_audit_run(run_path, project_root)
    expected = (
        expected_internal_artifact_paths(data)
        if layer.name == INTERNAL_ARTIFACT_DIR_NAME
        else expected_formal_artifact_paths(data)
    )
    expected_in_layer = {
        path.relative_to(layer.relative_to(project_root)).as_posix()
        for path in expected
        if path.is_relative_to(layer.relative_to(project_root))
    }
    actual_in_layer = enumerate_layer_files(layer)
    unexpected = sorted(actual_in_layer - expected_in_layer)

    # Do not resolve or open any other file once the filename-only check finds an extra artifact.
    if unexpected:
        return ArtifactAuditPlan(
            files=[run_path],
            audit_root=project_root,
            pre_audited_files={run_path},
            preflight_issues=preflight_issues,
            unexpected_artifacts=unexpected,
        )

    files = [run_path]
    for relative in sorted(actual_in_layer):
        candidate = layer / relative
        label = candidate.relative_to(project_root).as_posix()
        ensure_regular_project_file(candidate, project_root, label)
        if candidate != run_path:
            files.append(candidate)
    return ArtifactAuditPlan(
        files=files,
        audit_root=project_root,
        pre_audited_files={run_path},
        preflight_issues=preflight_issues,
        unexpected_artifacts=[],
    )


def explicit_file_audit_plan(target: Path) -> ArtifactAuditPlan:
    """Accept one directly named, locally auditable generated artifact only."""
    if not target.exists():
        raise FileNotFoundError(f"Artifact path not found: {target}")
    if not target.is_file() or target.is_symlink():
        raise ValidationError("artifact-audit requires one real explicit artifact file, not a symlink.")
    if target.suffix.lower() not in SUPPORTED_EXPLICIT_ARTIFACT_SUFFIXES:
        raise ValidationError(
            "artifact-audit does not support this explicit artifact type; use JSON, Markdown, DOCX, "
            "or a raster watermarked screenshot."
        )
    return ArtifactAuditPlan(
        files=[target.resolve()],
        audit_root=None,
        pre_audited_files=set(),
        preflight_issues=[],
        unexpected_artifacts=[],
    )


def artifact_audit_plan(target: Path) -> ArtifactAuditPlan:
    if target.is_file() or target.is_symlink():
        return explicit_file_audit_plan(target)
    return directory_audit_plan(target)


def artifact_audit_files(target: Path) -> list[Path]:
    """Return a closed audit set; reject a layer containing any extra filename."""
    plan = artifact_audit_plan(target)
    if plan.unexpected_artifacts:
        raise ValidationError(
            "unexpected artifact: " + ", ".join(plan.unexpected_artifacts)
        )
    return plan.files


def artifact_audit_path_issues(
    path: Path,
    label: str,
    *,
    run_data: dict[str, Any] | None = None,
) -> list[str]:
    """Inspect one already-approved generated artifact without external access."""
    suffix = path.suffix.lower()
    if suffix in TEXT_ARTIFACT_SUFFIXES:
        text = path.read_text(encoding="utf-8", errors="ignore")
        if suffix == ".json" and path.name == RUN_FILE_NAME:
            try:
                data = json.loads(text)
                if not isinstance(data, dict):
                    return [f"JSON root must be an object: {label}"]
                scan_sensitive_values(data)
            except (json.JSONDecodeError, ValidationError) as exc:
                return [str(exc)]
            return text_metadata_privacy_issues(text, label)
        return text_privacy_issues(text, label)
    if suffix == ".docx":
        issues: list[str] = []
        try:
            with zipfile.ZipFile(path) as archive:
                for member in archive.namelist():
                    lower_member = member.lower()
                    member_suffix = Path(lower_member).suffix
                    if lower_member.endswith((".xml", ".rels", ".txt", ".json")):
                        payload = archive.read(member)
                        if lower_member == "word/document.xml":
                            payload = replace_approved_docx_company_codes(
                                payload,
                                run_data,
                            )
                        text = payload.decode("utf-8", errors="ignore")
                        issues.extend(text_privacy_issues(text, f"{label}!{member}"))
                        if DOCX_WORD_TEXT_PART_PATTERN.fullmatch(lower_member):
                            issues.extend(
                                docx_text_container_privacy_issues(
                                    payload,
                                    f"{label}!{member}",
                                )
                            )
                    elif lower_member.startswith("word/media/") and member_suffix in RASTER_IMAGE_SUFFIXES:
                        issues.extend(
                            image_metadata_privacy_issues(
                                archive.read(member),
                                f"{label}!{member}!image-metadata",
                            )
                        )
                    elif lower_member.startswith("word/media/") and member_suffix == ".svg":
                        text = archive.read(member).decode("utf-8", errors="ignore")
                        issues.extend(text_privacy_issues(text, f"{label}!{member}"))
        except zipfile.BadZipFile:
            issues.append(f"invalid DOCX package: {label}")
        return issues
    if suffix in RASTER_IMAGE_SUFFIXES:
        return image_metadata_privacy_issues(
            path.read_bytes(),
            f"{label}!image-metadata",
        )
    raise ValidationError(f"Unsupported approved artifact type: {label}")


def template_check(path: Path) -> dict[str, Any]:
    if not path.is_file():
        raise FileNotFoundError(f"DOCX template not found: {path}")
    if path.suffix.lower() != ".docx":
        raise ValidationError("Template must use .docx extension.")

    issues: list[str] = []
    try:
        with zipfile.ZipFile(path) as archive:
            names = set(archive.namelist())
            if "word/document.xml" not in names:
                issues.append("missing word/document.xml")
            combined_xml: list[str] = []
            for name in sorted(names):
                lower_name = name.lower()
                suffix = Path(lower_name).suffix
                if lower_name.endswith((".xml", ".rels", ".txt")):
                    text = archive.read(name).decode("utf-8", errors="ignore")
                    combined_xml.append(text)
                    issues.extend(text_privacy_issues(text, name))
                elif lower_name.startswith("word/media/") and suffix in RASTER_IMAGE_SUFFIXES:
                    issues.extend(
                        image_metadata_privacy_issues(
                            archive.read(name),
                            f"{name}!image-metadata",
                        )
                    )
                elif lower_name.startswith("word/media/") and suffix == ".svg":
                    text = archive.read(name).decode("utf-8", errors="ignore")
                    issues.extend(text_privacy_issues(text, name))
            package_text = "\n".join(combined_xml)
            if not re.search(r"\bPAGE\b", package_text):
                issues.append("missing Word PAGE field")
    except zipfile.BadZipFile as exc:
        raise ValidationError(f"Template is not a valid DOCX package: {path}") from exc

    try:
        from docx import Document

        document = Document(str(path))
    except Exception as exc:
        raise ValidationError(f"Template cannot be opened as DOCX: {path}") from exc

    subject_marker_rows = []
    for table in document.tables:
        for row in table.rows:
            if "【SUBJECT_ROWS】" in "".join(cell.text for cell in row.cells):
                subject_marker_rows.append(row)
                if len(row.cells) < 3:
                    issues.append("subject marker table must contain at least three columns")
    if len(subject_marker_rows) != 1:
        issues.append("template must contain exactly one 【SUBJECT_ROWS】 marker row")

    query_markers = [
        paragraph for paragraph in document.paragraphs if "【QUERY_RESULT_ITEMS】" in paragraph.text
    ]
    if len(query_markers) != 1:
        issues.append("template must contain exactly one 【QUERY_RESULT_ITEMS】 marker paragraph")

    if issues:
        raise ValidationError("DOCX template check failed: " + "; ".join(dict.fromkeys(issues)))
    return {
        "path": str(path),
        "subject_marker_rows": 1,
        "query_marker_paragraphs": 1,
        "page_field": True,
        "privacy_safe": True,
    }


def artifact_audit(
    root: Path,
    *,
    run_data: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Scan a dedicated generated layer or one explicit file without external access."""
    plan = artifact_audit_plan(root)
    if run_data is None and plan.audit_root is not None:
        run_data = load_json(plan.audit_root / INTERNAL_ARTIFACT_DIR_NAME / RUN_FILE_NAME)
    if run_data is not None:
        validate_run(run_data, formal_mode="draft")
    allowed_company_credit_codes = (
        frozenset(company_credit_code_locations(run_data, "run").values())
        if run_data is not None
        else frozenset()
    )
    issues = list(plan.preflight_issues)
    if plan.unexpected_artifacts:
        issues.extend(f"unexpected artifact: {relative}" for relative in plan.unexpected_artifacts)
        return {
            "scanned_files": len(plan.pre_audited_files),
            "validated_company_credit_codes": len(allowed_company_credit_codes),
            "issues": sorted(dict.fromkeys(issues)),
            "ok": False,
            "limitations": [
                "Raster images are checked only for filenames and embedded metadata; no pixel OCR is performed.",
                "Visually confirm that every retained image does not display a full identity number.",
            ],
        }
    for path in plan.files:
        label = path.name if plan.audit_root is None else path.relative_to(plan.audit_root).as_posix()
        # A filename is never part of the verified run context: exact company
        # codes are allowed only in approved artifact contents, never names.
        issues.extend(text_privacy_issues(label, f"filename:{label}"))
        if path not in plan.pre_audited_files:
            issues.extend(
                artifact_audit_path_issues(
                    path,
                    label,
                    run_data=run_data,
                )
            )
    return {
        "scanned_files": len(plan.files),
        "validated_company_credit_codes": len(allowed_company_credit_codes),
        "issues": sorted(dict.fromkeys(issues)),
        "ok": not issues,
        "limitations": [
            "Raster images are checked only for filenames and embedded metadata; no pixel OCR is performed.",
            "Visually confirm that every retained image does not display a full identity number.",
        ],
    }


def normalized_hostname(value: Any, label: str) -> str:
    hostname = str(value or "").strip().lower().rstrip(".")
    if not hostname or any(character in hostname for character in " /\\:@"):
        raise ValidationError(f"{label} must be a hostname without credentials, path, port, or wildcard.")
    try:
        parsed = urlparse(f"https://{hostname}")
    except ValueError as exc:
        raise ValidationError(f"{label} must be a valid hostname.") from exc
    if parsed.hostname != hostname or parsed.path not in {"", "/"} or parsed.port is not None:
        raise ValidationError(f"{label} must be a plain hostname.")
    return hostname


def hostname_allowed(hostname: str, domains: list[str]) -> bool:
    host = hostname.lower().rstrip(".")
    return any(host == domain or host.endswith("." + domain) for domain in domains)


def validate_safe_resume_location(
    value: Any,
    label: str,
    *,
    allowed_domains: list[str],
    required: bool = False,
) -> str:
    text = str(value or "").strip()
    if not text:
        if required:
            raise ValidationError(f"{label} is required for a recoverable handoff.")
        return ""
    return validate_archival_url(
        text,
        label,
        allowed_domains=allowed_domains,
        allow_query=False,
    )


def validate_watermarked_capture(
    path: Path,
    label: str,
    *,
    evidence_id: str,
    queried_at: datetime,
    subject: str,
    source: str,
) -> None:
    """Validate nonblank pixels and watermark metadata without attempting OCR."""
    if path.suffix.lower() != ".png":
        raise ValidationError(f"{label} must be a PNG created by watermark_capture.py.")
    try:
        from PIL import Image, ImageChops

        with Image.open(path) as image:
            converted = image.convert("RGB")
            if converted.width < 2 or converted.height < 2:
                raise ValidationError(f"{label} is too small to support a result record.")
            corner = Image.new("RGB", converted.size, converted.getpixel((0, 0)))
            if ImageChops.difference(converted, corner).getbbox() is None:
                raise ValidationError(
                    f"{label} is blank or single-colour; pixel content cannot support the evidence gate."
                )
            metadata = dict(image.info)
        expected = {
            "EvidenceId": evidence_id,
            "Subject": subject,
            "Source": source,
            "CaptureKind": WATERMARKED_CAPTURE_KIND,
            "SourceContentVerified": "heuristic_pre_watermark_gate_passed",
        }
        for key, value in expected.items():
            if str(metadata.get(key) or "") != value:
                raise ValidationError(f"{label} has missing or mismatched {key} metadata.")
        metadata_time = require_timestamp(metadata.get("QueriedAt"), f"{label}.QueriedAt")
        if metadata_time != queried_at:
            raise ValidationError(f"{label} QueriedAt metadata does not match query_time.")
    except ValidationError:
        raise
    except Exception as exc:
        raise ValidationError(f"{label} must be a readable raster image.") from exc


def validate_attempt_trace(
    attempts: list[Any],
    *,
    label: str,
    browser_selection: str,
    execution_state: str,
    submit_state: str,
) -> dict[str, Any]:
    handoff_open = False
    resume_check_required = False
    submitted = False
    post_submit_unknown = False
    resubmit_confirmed = False
    result_region_loaded = False
    fallback_count = 0
    stopped = False
    events: list[str] = []
    last_time: datetime | None = None
    first_time: datetime | None = None
    last_submit_at: datetime | None = None
    last_load_at: datetime | None = None
    last_screenshot_at: datetime | None = None
    last_sensitive_review_at: datetime | None = None
    for index, raw in enumerate(attempts):
        attempt = require_object(raw, f"{label}[{index}]")
        reject_unknown_keys(attempt, {"sequence", "at", "event"}, f"{label}[{index}]")
        sequence = attempt.get("sequence")
        if not isinstance(sequence, int) or isinstance(sequence, bool) or sequence != index + 1:
            raise ValidationError(f"{label}[{index}].sequence must be the consecutive integer {index + 1}.")
        at = require_timestamp(attempt.get("at"), f"{label}[{index}].at")
        if last_time is not None and at < last_time:
            raise ValidationError(f"{label} timestamps must be nondecreasing.")
        if first_time is None:
            first_time = at
        last_time = at
        event = require_text(attempt.get("event"), f"{label}[{index}].event")
        if event not in ALLOWED_ATTEMPT_EVENTS:
            raise ValidationError(f"Unsupported attempt event: {event}")
        events.append(event)
        if stopped:
            raise ValidationError("No attempt event is allowed after user_stopped.")

        if handoff_open and event in HANDOFF_FORBIDDEN_EVENTS:
            raise ValidationError(
                f"{label}[{index}] performs a page operation while user login or verification is pending."
            )
        if resume_check_required and event not in {
            "resume_safety_check_passed",
            "login_required",
            "session_expired",
            "human_verification_required",
            "handoff_started",
        }:
            raise ValidationError(
                "No operation may follow user_resume_confirmed before resume_safety_check_passed."
            )
        if event == "browser_fallback":
            fallback_count += 1
            if browser_selection == "user_explicit":
                raise ValidationError("A user-explicit browser selection cannot fall back to another browser.")
            if submitted or handoff_open:
                raise ValidationError("Browser fallback is allowed only before the first query submission and handoff.")
            if fallback_count > 1:
                raise ValidationError("Automatic browser selection may fall back at most once.")
        if event in HANDOFF_START_EVENTS:
            if submitted:
                post_submit_unknown = True
                resubmit_confirmed = False
                result_region_loaded = False
                last_load_at = None
                last_screenshot_at = None
                last_sensitive_review_at = None
            handoff_open = True
            resume_check_required = False
        elif event == "user_resume_confirmed":
            if not handoff_open:
                raise ValidationError("user_resume_confirmed requires an open user handoff.")
            handoff_open = False
            resume_check_required = True
        elif event == "resume_safety_check_passed":
            if not resume_check_required or handoff_open:
                raise ValidationError(
                    "resume_safety_check_passed requires a user-confirmed handoff recovery."
                )
            resume_check_required = False
        elif event == "user_stopped":
            if not handoff_open:
                raise ValidationError("user_stopped requires an open user handoff.")
            handoff_open = False
            stopped = True
        elif event in {"network_error", "tool_error"} and submitted:
            post_submit_unknown = True
            resubmit_confirmed = False
            result_region_loaded = False
            last_load_at = None
            last_screenshot_at = None
            last_sensitive_review_at = None
        elif event == "explicit_resubmit_confirmed":
            if not post_submit_unknown or handoff_open or resume_check_required:
                raise ValidationError(
                    "explicit_resubmit_confirmed requires a safely resumed or technical post-submit unknown state."
                )
            resubmit_confirmed = True
        elif event == "query_submitted":
            if submitted and not post_submit_unknown:
                raise ValidationError("A repeated query submission requires a preceding unknown-submit state.")
            if post_submit_unknown and not resubmit_confirmed:
                raise ValidationError(
                    "A query cannot be resubmitted after an unknown submit state without explicit confirmation."
                )
            submitted = True
            result_region_loaded = False
            last_submit_at = at
            last_load_at = None
            last_screenshot_at = None
            last_sensitive_review_at = None
            if post_submit_unknown:
                post_submit_unknown = False
                resubmit_confirmed = False
        elif event == "result_region_loaded":
            if not submitted or post_submit_unknown:
                raise ValidationError("result_region_loaded requires the current query submission to be known.")
            result_region_loaded = True
            last_load_at = at
            last_screenshot_at = None
            last_sensitive_review_at = None
        elif event == "screenshot":
            if not result_region_loaded or last_load_at is None:
                raise ValidationError("screenshot requires a loaded result region in the current submission cycle.")
            last_screenshot_at = at
        elif event == "sensitive_result_review_confirmed":
            if not result_region_loaded or last_load_at is None:
                raise ValidationError(
                    "sensitive_result_review_confirmed requires a loaded result region."
                )
            last_sensitive_review_at = at

    if handoff_open and execution_state != "pending_user_action":
        raise ValidationError("An open user handoff must remain pending_user_action.")
    if execution_state == "pending_user_action" and not handoff_open:
        raise ValidationError("pending_user_action requires an unresolved login or verification handoff.")
    if resume_check_required:
        raise ValidationError("A resumed handoff requires resume_safety_check_passed before continuing.")
    if post_submit_unknown and submit_state != "unknown":
        raise ValidationError("A post-submit challenge or technical uncertainty requires submit_state=unknown.")
    if submit_state == "unknown" and not post_submit_unknown:
        raise ValidationError("submit_state=unknown requires an unresolved post-submit uncertainty.")
    if submit_state == "submitted" and not submitted:
        raise ValidationError("submit_state=submitted requires a query_submitted attempt.")
    if submit_state == "not_submitted" and submitted:
        raise ValidationError("submit_state=not_submitted conflicts with a query_submitted attempt.")
    if browser_selection == "automatic_primary" and fallback_count:
        raise ValidationError("automatic_primary cannot contain a browser fallback.")
    if browser_selection == "automatic_fallback" and fallback_count != 1:
        raise ValidationError("automatic_fallback requires exactly one browser fallback.")
    return {
        "query_submitted": submitted,
        "result_region_loaded": result_region_loaded,
        "events": events,
        "first_at": first_time,
        "last_event": events[-1],
        "last_submit_at": last_submit_at,
        "last_load_at": last_load_at,
        "last_screenshot_at": last_screenshot_at,
        "last_sensitive_review_at": last_sensitive_review_at,
    }


def validate_query_executions(
    data: dict[str, Any],
    scope_map: dict[str, dict[str, Any]],
    *,
    scope_confirmed_at: datetime,
) -> tuple[
    dict[str, dict[str, Any]],
    dict[str, dict[str, Any]],
    dict[str, dict[str, Any]],
]:
    executions = require_list(data.get("query_executions"), "query_executions")
    by_id: dict[str, dict[str, Any]] = {}
    by_scope: dict[str, dict[str, Any]] = {}
    trace_by_id: dict[str, dict[str, Any]] = {}
    for index, raw in enumerate(executions):
        execution = require_object(raw, f"query_executions[{index}]")
        reject_unknown_keys(
            execution,
            {
                "execution_id",
                "scope_item_id",
                "browser_surface",
                "browser_selection",
                "execution_state",
                "access_state",
                "submit_state",
                "block_reason",
                "attempts",
                "completion_evidence",
                "safe_resume_location",
            },
            f"query_executions[{index}]",
        )
        execution_id = require_text(
            execution.get("execution_id"), f"query_executions[{index}].execution_id"
        )
        if execution_id in by_id:
            raise ValidationError(f"Duplicate execution_id: {execution_id}")
        scope_item_id = require_text(
            execution.get("scope_item_id"), f"query_executions[{index}].scope_item_id"
        )
        if scope_item_id not in scope_map:
            raise ValidationError(f"Execution references unknown scope item: {execution_id}")
        if scope_item_id in by_scope:
            raise ValidationError(f"A scope item may have only one execution: {scope_item_id}")
        browser_surface = require_text(
            execution.get("browser_surface"), f"query_executions[{index}].browser_surface"
        )
        if browser_surface not in ALLOWED_BROWSER_SURFACES:
            raise ValidationError(f"Unsupported browser_surface: {browser_surface}")
        browser_selection = require_text(
            execution.get("browser_selection"), f"query_executions[{index}].browser_selection"
        )
        if browser_selection not in ALLOWED_BROWSER_SELECTIONS:
            raise ValidationError(f"Unsupported browser_selection: {browser_selection}")
        execution_state = require_text(
            execution.get("execution_state"), f"query_executions[{index}].execution_state"
        )
        if execution_state not in ALLOWED_EXECUTION_STATES:
            raise ValidationError(f"Unsupported execution_state: {execution_state}")
        access_state = require_text(
            execution.get("access_state"), f"query_executions[{index}].access_state"
        )
        if access_state not in ALLOWED_ACCESS_STATES:
            raise ValidationError(f"Unsupported access_state: {access_state}")
        submit_state = require_text(
            execution.get("submit_state"), f"query_executions[{index}].submit_state"
        )
        if submit_state not in ALLOWED_SUBMIT_STATES:
            raise ValidationError(f"Unsupported submit_state: {submit_state}")
        block_reason = str(execution.get("block_reason") or "").strip()
        if block_reason and block_reason not in ALLOWED_BLOCK_REASONS:
            raise ValidationError(f"Unsupported block_reason: {block_reason}")
        if execution_state in {"blocked", "failed"} and not block_reason:
            raise ValidationError(f"{execution_state} execution requires block_reason: {execution_id}")
        if execution_state in {"ready", "running", "completed"} and block_reason:
            raise ValidationError(f"{execution_state} execution must not retain block_reason: {execution_id}")
        if execution_state == "pending_user_action" and block_reason not in {
            "login_required",
            "session_expired",
            "human_verification_required",
            "submit_state_unknown",
        }:
            raise ValidationError("pending_user_action requires a recoverable authentication block_reason.")
        if execution_state == "blocked" and block_reason not in BLOCKED_REASONS:
            raise ValidationError(f"blocked execution has an invalid terminal reason: {execution_id}")
        if execution_state == "failed" and block_reason not in FAILED_REASONS:
            raise ValidationError(f"failed execution has an invalid terminal reason: {execution_id}")
        if block_reason in BLOCKED_REASONS and execution_state != "blocked":
            raise ValidationError(f"Blocked reason must map to execution_state=blocked: {execution_id}")
        if block_reason in FAILED_REASONS and execution_state != "failed":
            raise ValidationError(f"Technical reason must map to execution_state=failed: {execution_id}")
        if execution_state == "blocked" and block_reason != "user_stopped" and access_state != "limited":
            raise ValidationError(f"Access block requires access_state=limited: {execution_id}")
        if execution_state == "failed" and access_state != "limited":
            raise ValidationError(f"Technical failure requires access_state=limited: {execution_id}")
        if block_reason == "submit_state_unknown" and access_state not in {
            "login_required",
            "session_expired",
            "human_verification_required",
        }:
            raise ValidationError("submit_state_unknown requires an authentication or verification access state.")
        expected_access_for_reason = {
            "login_required": "login_required",
            "session_expired": "session_expired",
            "human_verification_required": "human_verification_required",
        }
        if block_reason in expected_access_for_reason and access_state != expected_access_for_reason[block_reason]:
            raise ValidationError(f"access_state does not match block_reason: {execution_id}")
        if execution_state == "completed" and access_state not in {"not_required", "authenticated"}:
            raise ValidationError("completed execution requires not_required or authenticated access.")

        attempts = require_list(execution.get("attempts"), f"query_executions[{index}].attempts")
        if not attempts:
            raise ValidationError(f"query_executions[{index}].attempts must not be empty.")
        trace = validate_attempt_trace(
            attempts,
            label=f"query_executions[{index}].attempts",
            browser_selection=browser_selection,
            execution_state=execution_state,
            submit_state=submit_state,
        )
        if trace["first_at"] < scope_confirmed_at:
            raise ValidationError(f"Execution predates query_scope confirmation: {execution_id}")
        if execution_state == "blocked" and trace["last_event"] != block_reason:
            raise ValidationError(f"blocked terminal reason must be the final attempt event: {execution_id}")
        if execution_state == "failed" and trace["last_event"] != block_reason:
            raise ValidationError(f"failed terminal reason must be the final attempt event: {execution_id}")
        if block_reason == "user_stopped" and "user_stopped" not in trace["events"]:
            raise ValidationError(f"user_stopped requires an explicit terminal event: {execution_id}")
        if block_reason and block_reason != "submit_state_unknown" and block_reason not in trace["events"]:
            raise ValidationError(f"block_reason is not supported by the attempt trace: {execution_id}")
        if block_reason == "submit_state_unknown" and not any(
            event in POST_SUBMIT_CHALLENGE_EVENTS for event in trace["events"]
        ):
            raise ValidationError(f"submit_state_unknown requires a post-submit challenge: {execution_id}")
        evidence = require_object(
            execution.get("completion_evidence"),
            f"query_executions[{index}].completion_evidence",
        )
        reject_unknown_keys(
            evidence,
            {
                "query_submitted",
                "result_region_loaded",
                "result_signal",
                "observed_at",
                "evidence_id",
                "capture_review",
            },
            f"query_executions[{index}].completion_evidence",
        )
        evidence_query_submitted = require_bool(
            evidence.get("query_submitted"),
            f"query_executions[{index}].completion_evidence.query_submitted",
        )
        evidence_region_loaded = require_bool(
            evidence.get("result_region_loaded"),
            f"query_executions[{index}].completion_evidence.result_region_loaded",
        )
        if evidence_query_submitted != trace["query_submitted"]:
            raise ValidationError(f"completion_evidence query_submitted conflicts with attempts: {execution_id}")
        if evidence_region_loaded != trace["result_region_loaded"]:
            raise ValidationError(f"completion_evidence result_region_loaded conflicts with attempts: {execution_id}")
        result_signal = require_text(
            evidence.get("result_signal"),
            f"query_executions[{index}].completion_evidence.result_signal",
        )
        if result_signal not in ALLOWED_RESULT_SIGNALS:
            raise ValidationError(f"Unsupported result_signal: {result_signal}")
        capture_review = require_text(
            evidence.get("capture_review"),
            f"query_executions[{index}].completion_evidence.capture_review",
        )
        if capture_review not in ALLOWED_CAPTURE_REVIEWS:
            raise ValidationError(f"Unsupported capture_review: {capture_review}")
        observed_at = str(evidence.get("observed_at") or "").strip()
        evidence_id = str(evidence.get("evidence_id") or "").strip()
        if result_signal == "not_observed":
            if observed_at or evidence_id or capture_review != "not_reviewed" or evidence_region_loaded:
                raise ValidationError(f"Unobserved completion evidence must remain empty: {execution_id}")
        else:
            observed_time = require_timestamp(
                observed_at, f"query_executions[{index}].completion_evidence.observed_at"
            )
            if observed_time < scope_confirmed_at:
                raise ValidationError(f"Completion evidence predates scope confirmation: {execution_id}")
            if trace["last_load_at"] is not None and observed_time < trace["last_load_at"]:
                raise ValidationError(f"Completion evidence predates the latest result load: {execution_id}")
            if execution_state != "completed":
                raise ValidationError(f"Only completed executions may retain a result signal: {execution_id}")
        if execution_state == "completed":
            if result_signal == "not_observed":
                raise ValidationError(f"completed execution requires an observed result signal: {execution_id}")
            if result_signal != "not_applicable" and (
                not evidence_query_submitted or not evidence_region_loaded or submit_state != "submitted"
            ):
                raise ValidationError(f"completed query result requires submitted and loaded evidence: {execution_id}")
        else:
            if result_signal != "not_observed":
                raise ValidationError(f"Incomplete execution cannot retain a legal result signal: {execution_id}")

        validate_safe_resume_location(
            execution.get("safe_resume_location"),
            f"query_executions[{index}].safe_resume_location",
            allowed_domains=scope_map[scope_item_id]["allowed_domains"],
            required=execution_state == "pending_user_action",
        )
        by_id[execution_id] = execution
        by_scope[scope_item_id] = execution
        trace_by_id[execution_id] = trace
    return by_id, by_scope, trace_by_id


def validate_query_scope(
    data: dict[str, Any],
    subject_map: dict[str, dict[str, Any]],
    *,
    allow_unconfirmed: bool = False,
) -> dict[str, dict[str, Any]]:
    scope = require_object(data.get("query_scope"), "query_scope")
    reject_unknown_keys(
        scope,
        {"status", "confirmed_at", "selection_mode", "items"},
        "query_scope",
    )
    status = require_text(scope.get("status"), "query_scope.status")
    if status not in {QUERY_SCOPE_STATUS, "pending_confirmation"}:
        raise ValidationError("query_scope.status must be user_confirmed or pending_confirmation.")
    if status != QUERY_SCOPE_STATUS and not allow_unconfirmed:
        raise ValidationError("query_scope.status must be user_confirmed before queries can run.")
    confirmed_at = str(scope.get("confirmed_at") or "").strip()
    if status == QUERY_SCOPE_STATUS:
        require_timestamp(confirmed_at, "query_scope.confirmed_at")
    selection_mode = require_text(scope.get("selection_mode"), "query_scope.selection_mode")
    if selection_mode not in {"default_profile", "custom", "default_plus_custom"}:
        raise ValidationError(f"Unsupported query_scope.selection_mode: {selection_mode}")
    items = require_list(scope.get("items"), "query_scope.items")
    if not items:
        if allow_unconfirmed and status != QUERY_SCOPE_STATUS:
            return {}
        raise ValidationError("query_scope.items must not be empty.")
    scope_map: dict[str, dict[str, Any]] = {}
    for index, raw in enumerate(items):
        item = require_object(raw, f"query_scope.items[{index}]")
        reject_unknown_keys(
            item,
            {
                "scope_item_id",
                "subject_id",
                "matter_category_id",
                "site_id",
                "site_name",
                "site_basis",
                "allowed_domains",
                "query_term_mode",
                "conditions",
            },
            f"query_scope.items[{index}]",
        )
        item_id = require_text(item.get("scope_item_id"), f"query_scope.items[{index}].scope_item_id")
        if item_id in scope_map:
            raise ValidationError(f"Duplicate scope_item_id: {item_id}")
        subject_id = require_text(item.get("subject_id"), f"query_scope.items[{index}].subject_id")
        if subject_id not in subject_map:
            raise ValidationError(f"Scope item references unknown subject: {item_id}")
        require_text(item.get("matter_category_id"), f"query_scope.items[{index}].matter_category_id")
        require_text(item.get("site_id"), f"query_scope.items[{index}].site_id")
        require_text(item.get("site_name"), f"query_scope.items[{index}].site_name")
        site_basis = require_text(item.get("site_basis"), f"query_scope.items[{index}].site_basis")
        if site_basis not in {"default_profile", "user_specified", "user_confirmed_suggestion"}:
            raise ValidationError(f"Unsupported site_basis: {site_basis}")
        if selection_mode == "default_profile" and site_basis != "default_profile":
            raise ValidationError("default_profile selection cannot include a custom site basis.")
        if selection_mode == "custom" and site_basis == "default_profile":
            raise ValidationError("custom selection cannot include a default-profile site.")
        domains = require_list(item.get("allowed_domains"), f"query_scope.items[{index}].allowed_domains")
        if not domains:
            raise ValidationError(f"query_scope.items[{index}].allowed_domains must not be empty.")
        normalized_domains = [
            normalized_hostname(domain, f"query_scope.items[{index}].allowed_domains[{domain_index}]")
            for domain_index, domain in enumerate(domains)
        ]
        query_term_mode = require_text(
            item.get("query_term_mode"),
            f"query_scope.items[{index}].query_term_mode",
        )
        if query_term_mode not in {"exact_subject_name", "company_credit_code", "user_manual_full_id"}:
            raise ValidationError(f"Unsupported query_term_mode: {query_term_mode}")
        matching_subject = subject_map[subject_id]
        if query_term_mode == "company_credit_code":
            if matching_subject["type"] != "company" or not str(matching_subject.get("credit_code") or "").strip():
                raise ValidationError(f"company_credit_code mode requires a company credit_code: {item_id}")
            validate_company_credit_code(matching_subject["credit_code"], f"subjects[{matching_subject['subject_id']}].credit_code")
        if query_term_mode == "user_manual_full_id" and matching_subject["type"] != "natural_person":
            raise ValidationError(f"user_manual_full_id is only for natural persons: {item_id}")
        conditions = require_list(item.get("conditions"), f"query_scope.items[{index}].conditions")
        if not conditions:
            raise ValidationError(f"query_scope.items[{index}].conditions must not be empty.")
        condition_ids: set[str] = set()
        for condition_index, raw_condition in enumerate(conditions):
            condition = require_object(raw_condition, f"query_scope.items[{index}].conditions[{condition_index}]")
            reject_unknown_keys(
                condition,
                {"condition_id", "field", "value"},
                f"query_scope.items[{index}].conditions[{condition_index}]",
            )
            condition_id = require_text(
                condition.get("condition_id"),
                f"query_scope.items[{index}].conditions[{condition_index}].condition_id",
            )
            if condition_id in condition_ids:
                raise ValidationError(f"Duplicate condition_id in scope item: {item_id}")
            condition_ids.add(condition_id)
            require_text(condition.get("field"), f"query_scope.items[{index}].conditions[{condition_index}].field")
            require_text(condition.get("value"), f"query_scope.items[{index}].conditions[{condition_index}].value")
        if not any(
            condition["field"] == "period" and str(condition["value"]).strip()
            for condition in conditions
        ):
            raise ValidationError(f"query_scope.items[{index}] must confirm a non-empty period condition.")
        normalized_item = copy.deepcopy(item)
        normalized_item["allowed_domains"] = normalized_domains
        scope_map[item_id] = normalized_item
    return scope_map


def scope_preview(data: dict[str, Any]) -> dict[str, Any]:
    """Return a privacy-safe subject/issue/site/condition matrix before browsing."""
    # Preview is itself a sensitive-data boundary. Validate the complete run
    # shape before exposing even a matrix, while permitting only the explicit
    # pending skeleton produced by prepare.
    scan_sensitive_values(data)
    if str(data.get("schema_version") or "") not in {"1.1", "1.2"}:
        raise ValidationError("scope-preview requires schema_version 1.1 or 1.2.")
    validate_run(data, formal_mode="none", allow_unconfirmed_scope=True)
    subjects = {
        subject["subject_id"]: subject
        for subject in require_list(data.get("subjects"), "subjects")
        if isinstance(subject, dict) and subject.get("subject_id")
    }
    scope = require_object(data.get("query_scope"), "query_scope")
    rows: list[dict[str, Any]] = []
    for raw in require_list(scope.get("items"), "query_scope.items"):
        item = require_object(raw, "query_scope.items[]")
        subject = subjects.get(str(item.get("subject_id") or ""))
        conditions = item.get("conditions") if isinstance(item.get("conditions"), list) else []
        rows.append(
            {
                "scope_item_id": item.get("scope_item_id", ""),
                "subject_id": item.get("subject_id", ""),
                "subject_name": subject["name"] if subject else "",
                "matter_category_id": item.get("matter_category_id", ""),
                "site_id": item.get("site_id", ""),
                "site_name": item.get("site_name", ""),
                "site_basis": item.get("site_basis", ""),
                "allowed_domains": list(item.get("allowed_domains") or []),
                "query_term_mode": item.get("query_term_mode", ""),
                "conditions": conditions,
            }
        )
    return {
        "status": scope.get("status", ""),
        "confirmed_at": scope.get("confirmed_at", ""),
        "selection_mode": scope.get("selection_mode", ""),
        "query_allowed": scope.get("status") == QUERY_SCOPE_STATUS,
        "rows": rows,
    }


def doctor_report(browser: str) -> dict[str, Any]:
    dependency_modules = {
        "python-docx": "docx",
        "Pillow": "PIL",
        "lxml": "lxml",
    }
    dependencies = {
        name: importlib.util.find_spec(module_name) is not None
        for name, module_name in dependency_modules.items()
    }
    template_ok = False
    template_message = "bundled template is missing"
    if BUNDLED_TEMPLATE.is_file() and all(dependencies.values()):
        try:
            template_check(BUNDLED_TEMPLATE)
            template_ok = True
            template_message = "bundled template passed structural and privacy checks"
        except (FileNotFoundError, OSError, ValidationError) as exc:
            template_message = str(exc)

    local_ready = sys.version_info >= (3, 10) and all(dependencies.values())
    if browser == "available" and local_ready:
        mode = "full-workflow"
    elif local_ready:
        mode = "local-processing-only"
    elif browser == "available":
        mode = "browser-only"
    else:
        mode = "planning-only"

    return {
        "python": {
            "version": platform.python_version(),
            "minimum": "3.10",
            "ok": sys.version_info >= (3, 10),
        },
        "dependencies": dependencies,
        "browser": browser,
        "libreoffice": shutil.which("soffice") or shutil.which("libreoffice") or "",
        "bundled_template": {
            "available": BUNDLED_TEMPLATE.is_file(),
            "ok": template_ok,
            "message": template_message,
        },
        "mode": mode,
    }


def validate_run(
    data: dict[str, Any],
    *,
    workpaper_root: Path | None = None,
    formal_mode: str = "draft",
    allow_unconfirmed_scope: bool = False,
) -> None:
    scan_sensitive_values(data)
    schema_version = str(data.get("schema_version") or "")
    if schema_version not in COMPATIBLE_SCHEMA_VERSIONS:
        raise ValidationError(f"Unsupported schema_version: {data.get('schema_version')}")
    if schema_version == "1.2":
        validate_persisted_text_safety(data)
        reject_unknown_keys(
            data,
            {
                "schema_version",
                "run",
                "subjects",
                "query_scope",
                "query_executions",
                "queries",
                "opinion_wording_requested",
            },
            "run root",
        )
    elif "query_executions" in data:
        raise ValidationError("query_executions is available only in schema_version 1.2.")

    run = require_object(data.get("run"), "run")
    if schema_version == "1.2":
        reject_unknown_keys(
            run,
            {"run_id", "timezone", "profile", "project", "formal_record"},
            "run",
        )
    require_text(run.get("run_id"), "run.run_id")
    timezone_name = require_text(run.get("timezone"), "run.timezone")
    try:
        ZoneInfo(timezone_name)
    except Exception as exc:  # pragma: no cover - platform timezone database edge
        raise ValidationError(f"Unknown timezone: {timezone_name}") from exc
    profile = require_text(run.get("profile"), "run.profile")
    if profile not in ALLOWED_PROFILES:
        raise ValidationError(f"run.profile must be one of: {', '.join(sorted(ALLOWED_PROFILES))}")

    project = require_object(run.get("project"), "run.project")
    if schema_version == "1.2":
        reject_unknown_keys(
            project,
            {"name", "short_name", "matter", "period"},
            "run.project",
        )
    require_text(project.get("name"), "run.project.name")
    require_text(project.get("short_name"), "run.project.short_name")
    require_text(project.get("matter"), "run.project.matter")
    require_text(project.get("period"), "run.project.period")

    formal_record = require_object(run.get("formal_record", {}), "run.formal_record")
    if schema_version == "1.2":
        reject_unknown_keys(
            formal_record,
            {"query_date", "query_location", "query_people"},
            "run.formal_record",
        )

    subjects = require_list(data.get("subjects"), "subjects")
    if not subjects:
        raise ValidationError("subjects must not be empty.")
    subject_map: dict[str, dict[str, Any]] = {}
    for index, raw in enumerate(subjects):
        subject = require_object(raw, f"subjects[{index}]")
        if schema_version == "1.2":
            reject_unknown_keys(
                subject,
                {
                    "subject_id",
                    "type",
                    "name",
                    "role",
                    "associated_entity",
                    "credit_code",
                    "formal_identifier_mode",
                    "masked_id_number",
                },
                f"subjects[{index}]",
            )
        subject_id = require_text(subject.get("subject_id"), f"subjects[{index}].subject_id")
        if subject_id in subject_map:
            raise ValidationError(f"Duplicate subject_id: {subject_id}")
        subject_type = require_text(subject.get("type"), f"subjects[{index}].type")
        if subject_type not in ALLOWED_SUBJECT_TYPES:
            raise ValidationError(f"Invalid subject type for {subject_id}: {subject_type}")
        require_text(subject.get("name"), f"subjects[{index}].name")
        require_text(subject.get("role"), f"subjects[{index}].role")
        identifier_mode = str(subject.get("formal_identifier_mode") or "user_fill").strip()
        if identifier_mode not in FORMAL_IDENTIFIER_MODES:
            raise ValidationError(f"Invalid formal_identifier_mode for {subject_id}: {identifier_mode}")
        if subject_type == "natural_person" and identifier_mode != "user_fill":
            raise ValidationError(f"Natural person formal_identifier_mode must be user_fill: {subject_id}")
        if subject_type == "company" and (
            identifier_mode == "auto_fill_company_credit_code"
            or str(subject.get("credit_code") or "").strip()
        ):
            validate_company_credit_code(
                subject.get("credit_code"),
                f"subjects[{index}].credit_code",
            )
        if subject_type == "natural_person" and str(subject.get("credit_code") or "").strip():
            raise ValidationError(f"Natural person must not use credit_code field: {subject_id}")
        masked_id = str(subject.get("masked_id_number") or "").strip()
        if subject_type == "company" and masked_id:
            raise ValidationError(f"Company must not use masked_id_number field: {subject_id}")
        if subject_type == "natural_person" and masked_id and not MASKED_ID_PATTERN.fullmatch(masked_id):
            raise ValidationError(
                "masked_id_number must keep 10 digits and mask the final 8 characters: " + subject_id
            )
        subject_map[subject_id] = subject

    scope_map: dict[str, dict[str, Any]] = {}
    if schema_version in {"1.1", "1.2"}:
        scope_map = validate_query_scope(
            data,
            subject_map,
            allow_unconfirmed=allow_unconfirmed_scope,
        )

    scope_confirmed_at: datetime | None = None
    if schema_version in {"1.1", "1.2"} and data.get("query_scope", {}).get("status") == QUERY_SCOPE_STATUS:
        scope_confirmed_at = require_timestamp(
            data["query_scope"].get("confirmed_at"), "query_scope.confirmed_at"
        )

    execution_by_id: dict[str, dict[str, Any]] = {}
    execution_by_scope: dict[str, dict[str, Any]] = {}
    trace_by_execution: dict[str, dict[str, Any]] = {}
    if schema_version == "1.2":
        if scope_confirmed_at is None and data.get("query_executions"):
            raise ValidationError("Schema 1.2 executions require a confirmed query scope.")
        if scope_confirmed_at is not None:
            execution_by_id, execution_by_scope, trace_by_execution = validate_query_executions(
                data,
                scope_map,
                scope_confirmed_at=scope_confirmed_at,
            )

    queries = require_list(data.get("queries"), "queries")
    if schema_version in {"1.1", "1.2"} and data.get("query_scope", {}).get("status") != QUERY_SCOPE_STATUS and queries:
        raise ValidationError("Unconfirmed query scope cannot contain queries.")
    evidence_ids: set[str] = set()
    query_by_execution: dict[str, dict[str, Any]] = {}
    for index, raw in enumerate(queries):
        query = require_object(raw, f"queries[{index}]")
        if schema_version == "1.2":
            reject_unknown_keys(
                query,
                {
                    "evidence_id",
                    "execution_id",
                    "scope_item_id",
                    "subject_id",
                    "matter_category_id",
                    "site_id",
                    "site_name",
                    "query_term_mode",
                    "condition_ids",
                    "conditions",
                    "url",
                    "query_time",
                    "query_terms",
                    "filters",
                    "status",
                    "result_summary",
                    "identity_assessment",
                    "follow_up",
                    "screenshot_path",
                    "capture_kind",
                },
                f"queries[{index}]",
            )
        evidence_id = require_text(query.get("evidence_id"), f"queries[{index}].evidence_id")
        if evidence_id in evidence_ids:
            raise ValidationError(f"Duplicate evidence_id: {evidence_id}")
        evidence_ids.add(evidence_id)
        subject_id = require_text(query.get("subject_id"), f"queries[{index}].subject_id")
        if subject_id not in subject_map:
            raise ValidationError(f"Query references unknown subject: {subject_id}")
        if schema_version in {"1.1", "1.2"}:
            scope_item_id = require_text(
                query.get("scope_item_id"),
                f"queries[{index}].scope_item_id",
            )
            if scope_item_id not in scope_map:
                raise ValidationError(f"Query references unknown scope item: {evidence_id}")
            scope_item = scope_map[scope_item_id]
            subject = subject_map[subject_id]
            if subject_id != scope_item["subject_id"]:
                raise ValidationError(f"Query subject exceeds confirmed scope: {evidence_id}")
            if require_text(query.get("matter_category_id"), f"queries[{index}].matter_category_id") != scope_item["matter_category_id"]:
                raise ValidationError(f"Query category exceeds confirmed scope: {evidence_id}")
            if require_text(query.get("site_id"), f"queries[{index}].site_id") != scope_item["site_id"]:
                raise ValidationError(f"Query site exceeds confirmed scope: {evidence_id}")
            if require_text(query.get("site_name"), f"queries[{index}].site_name") != scope_item["site_name"]:
                raise ValidationError(f"Query site name exceeds confirmed scope: {evidence_id}")
            query_term_mode = str(query.get("query_term_mode") or scope_item["query_term_mode"]).strip()
            if query_term_mode != scope_item["query_term_mode"]:
                raise ValidationError(f"Query term mode exceeds confirmed scope: {evidence_id}")
        if schema_version == "1.2":
            execution_id = require_text(
                query.get("execution_id"), f"queries[{index}].execution_id"
            )
            if execution_id not in execution_by_id:
                raise ValidationError(f"Query references unknown execution: {evidence_id}")
            if execution_id in query_by_execution:
                raise ValidationError(f"An execution may have only one legal result: {execution_id}")
            if execution_by_id[execution_id]["scope_item_id"] != scope_item_id:
                raise ValidationError(f"Query execution exceeds confirmed scope: {evidence_id}")
            query_by_execution[execution_id] = query
        require_text(query.get("site_id"), f"queries[{index}].site_id")
        require_text(query.get("site_name"), f"queries[{index}].site_name")
        url = require_text(query.get("url"), f"queries[{index}].url")
        if not re.match(r"^https?://", url, re.IGNORECASE):
            raise ValidationError(f"Query URL must use HTTP(S): {evidence_id}")
        if schema_version in {"1.1", "1.2"}:
            validate_archival_url(
                url,
                f"queries[{index}].url",
                allowed_domains=scope_item["allowed_domains"],
                allow_query=schema_version == "1.1",
            )
        query_time = require_timestamp(query.get("query_time"), f"queries[{index}].query_time")
        if schema_version in {"1.1", "1.2"} and scope_confirmed_at is not None and query_time < scope_confirmed_at:
            raise ValidationError(f"Query predates query_scope confirmation: {evidence_id}")
        if schema_version == "1.0":
            require_text(query.get("query_terms"), f"queries[{index}].query_terms")
            require_text(query.get("filters"), f"queries[{index}].filters")
        elif schema_version == "1.1":
            if "query_terms" in query or "filters" in query:
                raise ValidationError(
                    "schema 1.1 query_terms and filters are generated from confirmed structured scope."
                )
        else:
            require_text(query.get("query_terms"), f"queries[{index}].query_terms")
            require_text(query.get("filters"), f"queries[{index}].filters")
        if schema_version in {"1.1", "1.2"}:
            scope_conditions = scope_item["conditions"]
            query_condition_ids = require_list(
                query.get("condition_ids"),
                f"queries[{index}].condition_ids",
            )
            expected_condition_ids = [condition["condition_id"] for condition in scope_conditions]
            if query_condition_ids != expected_condition_ids:
                raise ValidationError(f"Query conditions exceed confirmed scope: {evidence_id}")
            query_conditions = require_list(query.get("conditions"), f"queries[{index}].conditions")
            if query_conditions != scope_conditions:
                raise ValidationError(f"Query condition values exceed confirmed scope: {evidence_id}")
            if schema_version == "1.2":
                expected_terms = canonical_query_terms(subject, scope_item["query_term_mode"])
                expected_filters = canonical_query_filters(scope_conditions)
                if query["query_terms"] != expected_terms:
                    raise ValidationError(f"Query terms exceed confirmed scope: {evidence_id}")
                if query["filters"] != expected_filters:
                    raise ValidationError(f"Query filters exceed confirmed scope: {evidence_id}")
        status = require_text(query.get("status"), f"queries[{index}].status")
        if status not in ALLOWED_STATUSES:
            raise ValidationError(f"Invalid status for {evidence_id}: {status}")
        summary = require_text(query.get("result_summary"), f"queries[{index}].result_summary")
        assessment = require_text(
            query.get("identity_assessment"), f"queries[{index}].identity_assessment"
        )
        check_wording(summary, f"queries[{index}].result_summary")
        check_wording(assessment, f"queries[{index}].identity_assessment")

        combined = summary + assessment
        formal_negative_wording = (
            "未发现与" in combined and "相符" in combined and "负面记录" in combined
        )
        if status == "no_match_displayed" and not (
            any(token in combined for token in ("未显示", "未发现可确认", "未检出匹配"))
            or formal_negative_wording
        ):
            raise ValidationError(
                "no_match_displayed must use limited no-match wording: " + evidence_id
            )
        if status in {"access_limited", "failed"} and any(
            token in combined for token in RESTRICTED_NEGATIVE_SIGNALS
        ):
            raise ValidationError(f"Restricted or failed query cannot claim no record: {evidence_id}")
        if schema_version == "1.2" and status == "access_limited" and (
            summary != CANONICAL_RESTRICTED_RESULT
            or assessment != CANONICAL_RESTRICTED_ASSESSMENT
            or str(query.get("follow_up") or "").strip() != CANONICAL_RESTRICTED_FOLLOW_UP
        ):
            raise ValidationError(
                f"access_limited must use the canonical non-negative result text: {evidence_id}"
            )
        if schema_version == "1.2" and status == "failed" and (
            summary != CANONICAL_FAILED_RESULT
            or assessment != CANONICAL_RESTRICTED_ASSESSMENT
            or str(query.get("follow_up") or "").strip() != CANONICAL_RESTRICTED_FOLLOW_UP
        ):
            raise ValidationError(
                f"failed must use the canonical non-negative result text: {evidence_id}"
            )

        screenshot_text = str(query.get("screenshot_path") or "").strip()
        capture_kind = str(query.get("capture_kind") or "").strip()
        sensitive_no_capture = capture_kind == SENSITIVE_NO_CAPTURE_KIND
        if status in SCREENSHOT_REQUIRED and not screenshot_text and not sensitive_no_capture:
            raise ValidationError(
                "Screenshot is required unless a sensitive page was not retained: " + evidence_id
            )
        if screenshot_text:
            screenshot_path = validate_safe_relative_path(
                screenshot_text, f"queries[{index}].screenshot_path"
            )
            if capture_kind != WATERMARKED_CAPTURE_KIND:
                raise ValidationError(f"Screenshot must be watermarked_page_only: {evidence_id}")
            if workpaper_root is not None:
                root = workpaper_root.resolve()
                candidate = (root / screenshot_path).resolve()
                try:
                    candidate.relative_to(root)
                except ValueError as exc:
                    raise ValidationError(f"Screenshot escapes workpaper root: {evidence_id}") from exc
                if not candidate.is_file():
                    raise FileNotFoundError(f"Screenshot not found: {candidate}")
                if schema_version == "1.2":
                    trace = trace_by_execution[execution_id]
                    if trace["last_screenshot_at"] is None:
                        raise ValidationError(
                            f"Watermarked evidence requires a screenshot attempt in the current cycle: {evidence_id}"
                        )
                    if query_time > trace["last_screenshot_at"]:
                        raise ValidationError(
                            f"query_time cannot follow the latest screenshot event: {evidence_id}"
                        )
                    validate_watermarked_capture(
                        candidate,
                        f"queries[{index}].screenshot_path",
                        evidence_id=evidence_id,
                        queried_at=query_time,
                        subject=subject["name"],
                        source=query["site_name"],
                    )
        elif sensitive_no_capture:
            if SENSITIVE_NO_CAPTURE_EXPLANATION not in (
                summary + str(query.get("follow_up") or "")
            ):
                raise ValidationError(
                    "Sensitive no-capture record must use the required protection explanation: "
                    + evidence_id
                )
            if schema_version == "1.2":
                execution = execution_by_id[execution_id]
                trace = trace_by_execution[execution_id]
                sensitive_observed_at = require_timestamp(
                    execution["completion_evidence"]["observed_at"],
                    f"query_executions[{execution_id}].completion_evidence.observed_at",
                )
                sensitive_review_at = trace["last_sensitive_review_at"]
                if not (
                    subject["type"] == "natural_person"
                    and query_term_mode == "user_manual_full_id"
                    and execution["completion_evidence"]["capture_review"] == "sensitive_no_capture"
                    and trace["last_load_at"] is not None
                    and sensitive_review_at is not None
                    and trace["last_load_at"] <= sensitive_observed_at <= sensitive_review_at
                ):
                    raise ValidationError(
                        "Sensitive no-capture is limited to a user-reviewed natural-person user-manual full-ID result: "
                        + evidence_id
                    )
        elif capture_kind:
            raise ValidationError(f"capture_kind is invalid without a screenshot: {evidence_id}")

    if schema_version == "1.2":
        signal_by_status = {
            "identity_match": {"matching_results_displayed"},
            "no_match_displayed": {"explicit_zero_count", "explicit_empty_state"},
            "same_name_candidates": {"same_name_candidates_displayed"},
            "not_applicable": {"not_applicable"},
            "pending_review": {
                "matching_results_displayed",
                "same_name_candidates_displayed",
            },
        }
        for execution_id, execution in execution_by_id.items():
            query = query_by_execution.get(execution_id)
            state = execution["execution_state"]
            if state in TERMINAL_EXECUTION_STATES and query is None:
                raise ValidationError(f"Terminal execution requires exactly one legal result: {execution_id}")
            if state not in TERMINAL_EXECUTION_STATES and query is not None:
                raise ValidationError(f"Incomplete execution cannot have a legal result: {execution_id}")
            if query is None:
                continue
            status = query["status"]
            if state == "blocked" and status != "access_limited":
                raise ValidationError(f"blocked execution must map to access_limited: {execution_id}")
            if state == "failed" and status != "failed":
                raise ValidationError(f"failed execution must map to failed: {execution_id}")
            if state == "completed" and status in {"access_limited", "failed"}:
                raise ValidationError(f"completed execution cannot map to a restricted result: {execution_id}")
            evidence = execution["completion_evidence"]
            if state == "completed":
                if evidence["evidence_id"] != query["evidence_id"]:
                    raise ValidationError(f"Completion evidence_id must match the query: {execution_id}")
                allowed_signals = signal_by_status.get(status)
                if allowed_signals is not None and evidence["result_signal"] not in allowed_signals:
                    raise ValidationError(f"Result signal does not match query status: {execution_id}")
                observed_time = require_timestamp(
                    evidence["observed_at"],
                    f"query_executions[{execution_id}].completion_evidence.observed_at",
                )
                query_time = require_timestamp(
                    query["query_time"], f"queries[{query['evidence_id']}].query_time"
                )
                if observed_time != query_time:
                    raise ValidationError(
                        f"query_time must match completion_evidence.observed_at: {execution_id}"
                    )
            if status == "no_match_displayed":
                if execution["access_state"] not in {"not_required", "authenticated"}:
                    raise ValidationError(f"No-match result requires usable page access: {execution_id}")
                if execution["submit_state"] != "submitted":
                    raise ValidationError(f"No-match result requires a known submitted query: {execution_id}")
                if not evidence["query_submitted"] or not evidence["result_region_loaded"]:
                    raise ValidationError(f"No-match result requires a loaded result region: {execution_id}")
                screenshot_text = str(query.get("screenshot_path") or "").strip()
                protected_no_capture = (
                    query.get("capture_kind") == SENSITIVE_NO_CAPTURE_KIND
                    and not screenshot_text
                    and evidence["capture_review"] == "sensitive_no_capture"
                    and subject_map[query["subject_id"]]["type"] == "natural_person"
                    and query["query_term_mode"] == "user_manual_full_id"
                    and trace_by_execution[execution_id]["last_sensitive_review_at"] is not None
                    and SENSITIVE_NO_CAPTURE_EXPLANATION
                    in (str(query.get("result_summary") or "") + str(query.get("follow_up") or ""))
                )
                page_capture = (
                    query.get("capture_kind") == WATERMARKED_CAPTURE_KIND
                    and bool(screenshot_text)
                    and evidence["capture_review"] == "conditions_and_result_visible"
                )
                if not page_capture and not protected_no_capture:
                    raise ValidationError(
                        f"No-match result requires reviewed page-only evidence or the fixed sensitive no-capture exception: {execution_id}"
                    )
                if page_capture:
                    trace = trace_by_execution[execution_id]
                    observed_time = require_timestamp(
                        evidence["observed_at"],
                        f"query_executions[{execution_id}].completion_evidence.observed_at",
                    )
                    if trace["last_load_at"] is None or trace["last_screenshot_at"] is None:
                        raise ValidationError(
                            f"No-match evidence requires load and screenshot events: {execution_id}"
                        )
                    if not (
                        trace["last_load_at"] <= observed_time <= trace["last_screenshot_at"]
                    ):
                        raise ValidationError(
                            f"No-match evidence times do not belong to the latest submission cycle: {execution_id}"
                        )
                if protected_no_capture:
                    review_time = trace_by_execution[execution_id]["last_sensitive_review_at"]
                    observed_time = require_timestamp(
                        evidence["observed_at"],
                        f"query_executions[{execution_id}].completion_evidence.observed_at",
                    )
                    if review_time is None or observed_time > review_time:
                        raise ValidationError(
                            f"Sensitive no-capture review must follow result observation: {execution_id}"
                        )

    if formal_mode not in {"none", "draft", "final"}:
        raise ValidationError("formal_mode must be none, draft, or final.")
    if formal_mode == "final":
        if schema_version == "1.2":
            if workpaper_root is None:
                raise ValidationError(
                    "Schema 1.2 final validation requires workpaper_root for evidence checks."
                )
            missing_scope = sorted(set(scope_map) - set(execution_by_scope))
            if missing_scope:
                raise ValidationError(
                    "Final output is blocked by unstarted scope item(s): " + ", ".join(missing_scope)
                )
            pending_scope = sorted(
                scope_item_id
                for scope_item_id, execution in execution_by_scope.items()
                if execution["execution_state"] not in TERMINAL_EXECUTION_STATES
                or execution["submit_state"] == "unknown"
            )
            if pending_scope:
                raise ValidationError(
                    "Final output is blocked by incomplete or unknown-submit scope item(s): "
                    + ", ".join(pending_scope)
                )
        formal = require_object(run.get("formal_record"), "run.formal_record")
        require_text(formal.get("query_date"), "run.formal_record.query_date")
        require_text(formal.get("query_location"), "run.formal_record.query_location")
        people = require_list(formal.get("query_people"), "run.formal_record.query_people")
        if not people or any(not str(item).strip() for item in people):
            raise ValidationError("run.formal_record.query_people must not be empty in final mode.")
        for subject in subjects:
            if subject["type"] == "company":
                if str(subject.get("credit_code") or "").strip():
                    validate_company_credit_code(
                        subject["credit_code"],
                        f"subjects[{subject['subject_id']}].credit_code",
                    )


def manual_identifier_fields_pending(data: dict[str, Any]) -> list[str]:
    pending: list[str] = []
    for subject in require_list(data.get("subjects"), "subjects"):
        mode = str(subject.get("formal_identifier_mode") or "user_fill").strip()
        if subject["type"] == "company" and mode == "auto_fill_company_credit_code":
            continue
        if subject["type"] == "company":
            pending.append(f"{subject['subject_id']}.formal_identifier_user_fill")
        else:
            pending.append(f"{subject['subject_id']}.formal_identifier_user_fill")
    return pending


def completion_preview(
    data: dict[str, Any], *, workpaper_root: Path | None = None
) -> dict[str, Any]:
    """Return a read-only execution coverage summary for a schema 1.2 run."""
    if str(data.get("schema_version") or "") != "1.2":
        raise ValidationError("completion-preview requires schema_version 1.2.")
    validate_run(
        data,
        workpaper_root=None,
        formal_mode="draft",
        allow_unconfirmed_scope=True,
    )
    scope = require_object(data.get("query_scope"), "query_scope")
    items = require_list(scope.get("items"), "query_scope.items")
    executions = {
        item["scope_item_id"]: item
        for item in require_list(data.get("query_executions"), "query_executions")
    }
    counts = {
        "not_started": 0,
        "pending_user_action": 0,
        "ready": 0,
        "running": 0,
        "completed": 0,
        "blocked": 0,
        "failed": 0,
    }
    blocking_items: list[dict[str, Any]] = []
    for item in items:
        scope_item_id = item["scope_item_id"]
        execution = executions.get(scope_item_id)
        if execution is None:
            counts["not_started"] += 1
            blocking_items.append(
                {
                    "scope_item_id": scope_item_id,
                    "execution_id": "",
                    "execution_state": "not_started",
                    "access_state": "unknown",
                    "submit_state": "not_submitted",
                    "block_reason": "",
                }
            )
            continue
        state = execution["execution_state"]
        counts[state] += 1
        if state not in TERMINAL_EXECUTION_STATES or execution["submit_state"] == "unknown":
            blocking_items.append(
                {
                    "scope_item_id": scope_item_id,
                    "execution_id": execution["execution_id"],
                    "execution_state": state,
                    "access_state": execution["access_state"],
                    "submit_state": execution["submit_state"],
                    "block_reason": execution.get("block_reason", ""),
                }
            )
    execution_ready = (
        scope.get("status") == QUERY_SCOPE_STATUS
        and len(items) > 0
        and not blocking_items
        and sum(counts[state] for state in TERMINAL_EXECUTION_STATES) == len(items)
    )
    readiness_issues: list[str] = []
    final_ready = False
    if not execution_ready:
        readiness_issues.append("execution_scope_incomplete")
    elif workpaper_root is None:
        readiness_issues.append("workpaper_root_required")
    else:
        try:
            validate_run(
                data,
                workpaper_root=workpaper_root,
                formal_mode="final",
                allow_unconfirmed_scope=False,
            )
            final_ready = True
        except (FileNotFoundError, OSError, ValidationError) as exc:
            readiness_issues.append(str(exc))
    return {
        "schema_version": "1.2",
        "scope_status": scope.get("status", ""),
        "total_scope_items": len(items),
        "counts": counts,
        "blocking_items": blocking_items,
        "execution_ready": execution_ready,
        "readiness_issues": readiness_issues,
        "final_ready": final_ready,
    }


def execution_completion_key(
    execution: dict[str, Any] | None,
    query: dict[str, Any] | None,
) -> str:
    if execution is None:
        return "not_started"
    if execution["submit_state"] == "unknown":
        return "submit_unknown"
    state = execution["execution_state"]
    if state == "pending_user_action":
        if execution["access_state"] == "human_verification_required":
            return "pending_verification"
        return "pending_login"
    if state in {"ready", "running"}:
        return "in_progress"
    if state == "blocked":
        return "access_limited"
    if state == "failed":
        return "failed"
    if query is None:
        return "in_progress"
    return {
        "identity_match": "completed_match",
        "no_match_displayed": "completed_no_match",
        "same_name_candidates": "same_name",
        "not_applicable": "not_applicable",
        "access_limited": "access_limited",
        "failed": "failed",
        "pending_review": "pending_review",
    }[query["status"]]


def canonical_query_terms(subject: dict[str, Any], query_term_mode: str) -> str:
    if query_term_mode == "exact_subject_name":
        return f"完整主体名称：{subject['name']}"
    if query_term_mode == "company_credit_code":
        return f"完整主体名称：{subject['name']}；统一社会信用代码：引用 subjects[].credit_code（不在查询描述重复）"
    if query_term_mode == "user_manual_full_id":
        return "完整身份证号码：由用户本人在目标网站手工输入（实际值不写入运行文件）"
    raise ValidationError(f"Unsupported query_term_mode: {query_term_mode}")


def canonical_query_filters(conditions: list[dict[str, Any]]) -> str:
    return "；".join(f"{condition['field']}：{condition['value']}" for condition in conditions)


def query_display_descriptions(
    schema_version: str,
    subject: dict[str, Any],
    query: dict[str, Any],
) -> tuple[str, str]:
    """Return privacy-safe deterministic query descriptions for workpapers."""
    if schema_version == "1.0":
        return (
            require_text(query.get("query_terms"), "query.query_terms"),
            require_text(query.get("filters"), "query.filters"),
        )
    term_mode = require_text(query.get("query_term_mode"), "query.query_term_mode")
    if term_mode == "exact_subject_name":
        query_terms = "完整公司名称" if subject["type"] == "company" else "完整自然人姓名"
    elif term_mode == "company_credit_code":
        query_terms = "经用户确认的统一社会信用代码"
    elif term_mode == "user_manual_full_id":
        query_terms = "由用户本人在目标网站手工输入的完整身份证号码（未由 Agent 或 Skill 读取或保存）"
    else:
        raise ValidationError(f"Unsupported query_term_mode: {term_mode}")
    conditions = require_list(query.get("conditions"), "query.conditions")
    filters = "；".join(
        f"{require_text(condition.get('field'), 'query.condition.field')}："
        f"{require_text(condition.get('value'), 'query.condition.value')}"
        for condition in conditions
        if isinstance(condition, dict)
    )
    if not filters or len(conditions) != sum(isinstance(item, dict) for item in conditions):
        raise ValidationError("query.conditions must contain only structured condition objects.")
    return query_terms, filters


def prepare_run(input_data: dict[str, Any]) -> dict[str, Any]:
    scan_sensitive_values(input_data, "input")
    requested_schema = str(input_data.get("schema_version") or SCHEMA_VERSION)
    if requested_schema not in COMPATIBLE_SCHEMA_VERSIONS:
        raise ValidationError(f"Unsupported schema_version: {requested_schema}")
    if requested_schema != SCHEMA_VERSION:
        raise ValidationError(
            "prepare only creates schema 1.2; schema 1.0 and 1.1 are validate/build read compatibility only."
        )
    project = require_object(input_data.get("project"), "project")
    subjects = copy.deepcopy(require_list(input_data.get("subjects"), "subjects"))
    formal_input = require_object(input_data.get("formal_record", {}), "formal_record")
    now = datetime.now(ZoneInfo(DEFAULT_TIMEZONE))
    prepared = {
        "schema_version": requested_schema,
        "run": {
            "run_id": str(input_data.get("run_id") or f"NQ-{now:%Y%m%d-%H%M%S}"),
            "timezone": str(input_data.get("timezone") or DEFAULT_TIMEZONE),
            "profile": str(input_data.get("profile") or "custom"),
            "project": {
                "name": require_text(project.get("name"), "project.name"),
                "short_name": require_text(project.get("short_name"), "project.short_name"),
                "matter": require_text(project.get("matter"), "project.matter"),
                "period": require_text(project.get("period"), "project.period"),
            },
            "formal_record": {
                "query_date": str(formal_input.get("query_date") or ""),
                "query_location": str(formal_input.get("query_location") or ""),
                "query_people": copy.deepcopy(formal_input.get("query_people", [])),
            },
        },
        "subjects": subjects,
        "queries": copy.deepcopy(input_data.get("queries", [])),
        "query_executions": copy.deepcopy(input_data.get("query_executions", [])),
        "opinion_wording_requested": bool(input_data.get("opinion_wording_requested", False)),
    }
    prepared["query_scope"] = copy.deepcopy(
        input_data.get(
            "query_scope",
            {
                "status": "pending_confirmation",
                "confirmed_at": "",
                "selection_mode": "custom",
                "items": [],
            },
        )
    )
    validate_run(prepared, formal_mode="draft", allow_unconfirmed_scope=True)
    return prepared


def safe_filename(value: str) -> str:
    text = re.sub(r"[\\/:*?\"<>|]", "-", value.strip())
    text = re.sub(r"\s+", "", text)
    text = text.strip(".-")
    return text or "未命名"


def markdown_cell(value: Any) -> str:
    return str(value or "").replace("|", "\\|").replace("\r", " ").replace("\n", " ").strip()


def markdown_screenshot_link(
    screenshot_path: str,
    *,
    workpaper_root: Path | None = None,
    markdown_dir: Path | None = None,
) -> str:
    if workpaper_root is None or markdown_dir is None:
        return Path(screenshot_path).as_posix()
    source = (workpaper_root.resolve() / Path(screenshot_path)).resolve()
    return Path(os.path.relpath(source, markdown_dir.resolve())).as_posix()


def query_date_range(queries: list[dict[str, Any]], timezone_name: str) -> str:
    if not queries:
        return "尚未查询"
    timezone = ZoneInfo(timezone_name)
    dates = sorted(
        require_timestamp(item["query_time"], "query.query_time").astimezone(timezone).date()
        for item in queries
    )
    if dates[0] == dates[-1]:
        return f"{dates[0].year}年{dates[0].month}月{dates[0].day}日（北京时间）"
    return (
        f"{dates[0].year}年{dates[0].month}月{dates[0].day}日至"
        f"{dates[-1].year}年{dates[-1].month}月{dates[-1].day}日（北京时间）"
    )


def build_opinion_wording(subject: dict[str, Any], queries: list[dict[str, Any]]) -> str:
    source_names = "、".join(dict.fromkeys(item["site_name"] for item in queries))
    matched = [item for item in queries if item["status"] == "identity_match"]
    same_name = [item for item in queries if item["status"] == "same_name_candidates"]
    limited = [item for item in queries if item["status"] in {"access_limited", "failed"}]
    parts = [f"本所通过{source_names or '相关公开网站'}对{subject['name']}进行了公开信息核查。"]
    if matched:
        parts.append("公开页面显示与姓名及关联机构等身份要素相匹配的信息。")
    if same_name:
        parts.append("部分网站存在同名候选结果，现有公开信息不足以确认归属于本案核查对象。")
    if limited:
        parts.append("部分查询因网站功能或访问条件受限，相关事项仍需补充核验。")
    parts.append(
        "前述结果仅反映本次查询条件及网站公开可查询范围，不等同于对不存在任何诉讼、执行、处罚、失信或其他负面事项的绝对确认。"
    )
    return "".join(parts)


def sentence_text(value: Any) -> str:
    text = str(value or "").strip().rstrip("。；; ")
    return f"{text}。" if text else ""


def internal_conclusion_text(query: dict[str, Any], subject_name: str) -> str:
    summary = sentence_text(query.get("result_summary"))
    status = query["status"]
    if status == "identity_match":
        return summary + "经本次公开页面所示身份要素核对，相关结果与本案核查对象相符。"
    if status == "no_match_displayed":
        return (
            summary
            + f"仅就本次查询条件及该网站公开可查询范围，未显示与{subject_name}相符的公开负面记录。"
        )
    if status == "same_name_candidates":
        return summary + "相关同名记录尚需进一步核对，暂不归属于本案核查对象。"
    if status == "not_applicable":
        return summary + "该公开入口不适用于本次核查对象，本次未据此作出否定性结论。"
    if status == "access_limited":
        return summary + "该项查询未完整完成，不作为“未发现负面记录”的依据。"
    if status == "failed":
        return summary + "该项查询未完成，不作为“未发现负面记录”的依据。"
    return summary + "相关结果尚待进一步核对，暂不作结论。"


def build_internal_markdown(
    data: dict[str, Any],
    subject: dict[str, Any],
    queries: list[dict[str, Any]],
    *,
    workpaper_root: Path | None = None,
    markdown_dir: Path | None = None,
) -> str:
    run = data["run"]
    project = run["project"]
    lines: list[str] = [
        f"# {project['name']}",
        "",
        f"## {subject['name']}网络核查底稿",
        "",
        f"**核查日期：** {query_date_range(queries, run['timezone'])}  ",
        f"**核查对象：** {subject['name']}  ",
        f"**主体角色：** {subject['role']}  ",
    ]
    associated = str(subject.get("associated_entity") or "").strip()
    if associated:
        lines.append(f"**关联机构：** {associated}  ")
    lines.extend(
        [
            f"**核查目的：** {project['matter']}  ",
            "**底稿性质：** 公开信息核查底稿，不替代身份要素、主管机关文件、当事人声明及其他书面核验，也不单独构成不存在负面事项的最终结论。",
        ]
    )
    is_schema_12 = str(data.get("schema_version") or "") == "1.2"
    if is_schema_12:
        scope_items = [
            item
            for item in data["query_scope"]["items"]
            if item["subject_id"] == subject["subject_id"]
        ]
        executions_by_scope = {
            item["scope_item_id"]: item for item in data.get("query_executions", [])
        }
        queries_by_execution = {
            item["execution_id"]: item
            for item in queries
            if str(item.get("execution_id") or "")
        }
        lines.extend(
            [
                "",
                "## 一、范围项完成度",
                "",
                "| 范围项 | 核查事项 | 网站 | 完成状态 | 证据编号 |",
                "| --- | --- | --- | --- | --- |",
            ]
        )
        for item in scope_items:
            execution = executions_by_scope.get(item["scope_item_id"])
            query = queries_by_execution.get(execution["execution_id"]) if execution else None
            completion_key = execution_completion_key(execution, query)
            evidence_id = query["evidence_id"] if query else "—"
            lines.append(
                "| {scope} | {matter} | {site} | {status} | {evidence} |".format(
                    scope=markdown_cell(item["scope_item_id"]),
                    matter=markdown_cell(item["matter_category_id"]),
                    site=markdown_cell(item["site_name"]),
                    status=EXECUTION_STATUS_LABELS[completion_key],
                    evidence=evidence_id,
                )
            )
        if not scope_items:
            lines.append("| — | — | — | 未开始 | — |")
    mouth_index = "二" if is_schema_12 else "一"
    table_index = "三" if is_schema_12 else "二"
    detail_index = "四" if is_schema_12 else "三"
    conclusion_index = "五" if is_schema_12 else "四"
    followup_index = "六" if is_schema_12 else "五"
    opinion_index = "七" if is_schema_12 else "六"
    lines.extend(
        [
            "",
            f"## {mouth_index}、核查口径",
            "",
            f"1. 本次核查期间为{project['period']}，实际查询日期见各网站记录。",
            "2. 仅名称或姓名相同的结果列为同名候选，不直接归属于本案核查对象。",
            "3. 使用身份交叉条件时只记录‘本案材料所载身份要素’等非敏感描述，不留存完整身份证号码、出生日期等敏感值。",
            "4. 查询受限、失败或网站不适用的，不作为未发现记录的依据。",
            "5. 网络公开结果可能存在滞后、匿名化、未公开、下架或未收录等限制，应结合其他书面材料综合判断。",
            "",
            f"## {table_index}、核查结果总表",
            "",
            "| 序号 | 网站 | 当前查询入口 | 查询结果 | 身份判断 | 结论级别 | 截图编号 |",
            "| --- | --- | --- | --- | --- | --- | --- |",
        ]
    )
    for index, query in enumerate(queries, start=1):
        lines.append(
            "| {index} | {site} | <{url}> | {summary} | {assessment} | {status} | {evidence} |".format(
                index=index,
                site=markdown_cell(query["site_name"]),
                url=query["url"],
                summary=markdown_cell(query["result_summary"]),
                assessment=markdown_cell(query["identity_assessment"]),
                status=STATUS_LABELS[query["status"]],
                evidence=query["evidence_id"],
            )
        )

    lines.extend(["", f"## {detail_index}、逐站核查记录及截图", ""])
    for index, query in enumerate(queries, start=1):
        display_terms, display_filters = query_display_descriptions(
            str(data.get("schema_version") or ""), subject, query
        )
        lines.extend(
            [
                f"### {index}. {query['site_name']}",
                "",
                f"- 查询入口：<{query['url']}>",
                f"- 查询条件：{display_terms}",
                f"- 筛选条件：{display_filters}",
                f"- 查询结果：{query['result_summary']}",
                f"- 身份判断：{query['identity_assessment']}",
                f"- 结论级别：{STATUS_LABELS[query['status']]}",
            ]
        )
        screenshot = str(query.get("screenshot_path") or "").strip()
        if screenshot:
            link = markdown_screenshot_link(
                screenshot,
                workpaper_root=workpaper_root,
                markdown_dir=markdown_dir,
            )
            lines.append(f"- [查询截图（{query['evidence_id']}）]({link})")
        else:
            lines.append("- 截图：未取得；原因见查询结果和身份判断。")
        follow_up = str(query.get("follow_up") or "").strip()
        if follow_up:
            lines.append(f"- 后续核验：{follow_up}")
        lines.append("")

    lines.extend([f"## {conclusion_index}、初步结论", ""])
    if not queries:
        lines.append("尚未完成公开网站查询。")
    else:
        for index, query in enumerate(queries, start=1):
            lines.append(
                f"{index}. **{query['site_name']}：** "
                f"{internal_conclusion_text(query, subject['name'])}"
            )

    followups = list(
        dict.fromkeys(str(item.get("follow_up") or "").strip() for item in queries if str(item.get("follow_up") or "").strip())
    )
    lines.extend(["", f"## {followup_index}、建议补核事项", ""])
    if followups:
        for index, item in enumerate(followups, start=1):
            lines.append(f"{index}. {item}")
    else:
        lines.append("无。")

    if data.get("opinion_wording_requested"):
        lines.extend(
            [
                "",
                f"## {opinion_index}、可供成果文件使用的核查表述（待律师复核）",
                "",
                f"> {build_opinion_wording(subject, queries)}",
            ]
        )
    lines.append("")
    return "\n".join(lines)


def iter_paragraphs(document: Any):
    for paragraph in document.paragraphs:
        yield paragraph
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph


def replace_in_paragraph(paragraph: Any, replacements: dict[str, str]) -> None:
    for run in paragraph.runs:
        for old, new in replacements.items():
            if old in run.text:
                run.text = run.text.replace(old, new)
    combined = paragraph.text
    if not any(old in combined for old in replacements):
        return
    for old, new in replacements.items():
        combined = combined.replace(old, new)
    if paragraph.runs:
        paragraph.runs[0].text = combined
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(combined)


def set_run_font(run: Any, size_pt: float = 12, bold: bool | None = None) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt

    run.font.name = OUTPUT_DOCX_FONT
    run.font.size = Pt(size_pt)
    if bold is not None:
        run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), OUTPUT_DOCX_FONT)
    rfonts.set(qn("w:ascii"), OUTPUT_DOCX_FONT)
    rfonts.set(qn("w:hAnsi"), OUTPUT_DOCX_FONT)
    rfonts.set(qn("w:cs"), OUTPUT_DOCX_FONT)


def set_cell_text(cell: Any, value: str, *, bold: bool = False) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

    cell.text = value
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = None
        paragraph.paragraph_format.left_indent = None
        paragraph.paragraph_format.right_indent = None
        for run in paragraph.runs:
            set_run_font(run, 10.5, bold=bold)


def formal_identifier_text(subject: dict[str, Any], formal_mode: str) -> str:
    identifier_mode = str(subject.get("formal_identifier_mode") or "user_fill").strip()
    if subject["type"] == "company" and identifier_mode == "auto_fill_company_credit_code":
        return str(subject.get("credit_code") or "").strip()
    return "【待用户填写】" if formal_mode == "draft" else ""


def clear_paragraph_numbering(paragraph: Any) -> None:
    from docx.oxml.ns import qn

    ppr = paragraph._p.get_or_add_pPr()
    numpr = ppr.find(qn("w:numPr"))
    if numpr is not None:
        ppr.remove(numpr)


def concise_formal_query_parts(
    index: int,
    query: dict[str, Any],
    subject_name: str,
    *,
    query_terms: str | None = None,
    filters: str | None = None,
) -> tuple[str, str, str]:
    summary = str(query["result_summary"]).rstrip("。；; ")
    assessment = str(query["identity_assessment"]).rstrip("。；; ")
    status = query["status"]

    if status == "identity_match":
        result_text = f"{summary}。经核对，{assessment}。"
    elif status == "no_match_displayed":
        result_text = (
            f"截至查询日，经查询{query['site_name']}，{summary}。"
            f"前述结果仅反映本次查询条件及该网站公开可查询范围，"
            f"未显示与{subject_name}相符的公开负面记录。"
        )
    elif status == "same_name_candidates":
        result_text = (
            f"检索结果中存在同名记录，相关记录尚需进一步核对，"
            f"暂不作是否与{subject_name}相关的结论。"
        )
    elif status == "not_applicable":
        result_text = f"{summary}。本项未据此作出是否存在负面记录的结论。"
    elif status == "access_limited":
        result_text = "因网站访问、验证或权限条件限制，本项查询未完整完成，暂不作是否存在负面记录的结论。"
    elif status == "failed":
        result_text = "因技术原因，本项查询未完成，暂不作是否存在负面记录的结论。"
    else:
        result_text = f"{summary}。相关结果尚待进一步核对，暂不作结论。"

    displayed_terms = query_terms or require_text(query.get("query_terms"), "query.query_terms")
    displayed_filters = filters or require_text(query.get("filters"), "query.filters")
    return (
        f"{index}. 就{subject_name}查询{query['site_name']}（",
        str(query["url"]),
        f"），本次以{displayed_terms}为主要条件，并结合{displayed_filters}进行核查。{result_text}",
    )


def concise_formal_query_text(index: int, query: dict[str, Any], subject_name: str) -> str:
    prefix, url, suffix = concise_formal_query_parts(index, query, subject_name)
    return prefix + url + suffix


def draft_incomplete_scope_text(
    index: int,
    scope_item: dict[str, Any],
    execution: dict[str, Any] | None,
    subject_name: str,
) -> str:
    if execution is None:
        state_text = "本项尚未开始"
    elif execution["submit_state"] == "unknown":
        state_text = "查询提交状态尚待确认"
    elif execution["access_state"] in {"login_required", "session_expired"}:
        state_text = "待用户完成登录，本项尚未完成"
    elif execution["access_state"] == "human_verification_required":
        state_text = "待用户完成人工验证，本项尚未完成"
    else:
        state_text = "本项尚未完成"
    return (
        f"{index}. 就{subject_name}查询{scope_item['site_name']}，{state_text}，"
        "暂不作是否存在相关记录的结论。"
    )


def breakable_url_display(url: str) -> str:
    display = re.sub(r"%[0-9A-Fa-f]{2}", lambda match: match.group(0) + ZERO_WIDTH_SPACE, url)
    return re.sub(r"[/\\.?&=]", lambda match: match.group(0) + ZERO_WIDTH_SPACE, display)


def enable_latin_word_wrap(paragraph: Any) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    ppr = paragraph._p.get_or_add_pPr()
    existing = ppr.find(qn("w:wordWrap"))
    if existing is not None:
        existing.set(qn("w:val"), "1")
        return
    word_wrap = OxmlElement("w:wordWrap")
    word_wrap.set(qn("w:val"), "1")
    ppr.append(word_wrap)


def clear_paragraph_content(paragraph: Any) -> None:
    from docx.oxml.ns import qn

    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def append_external_hyperlink(paragraph: Any, url: str, display_text: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), OUTPUT_DOCX_FONT)
    fonts.set(qn("w:hAnsi"), OUTPUT_DOCX_FONT)
    fonts.set(qn("w:eastAsia"), OUTPUT_DOCX_FONT)
    fonts.set(qn("w:cs"), OUTPUT_DOCX_FONT)
    run_properties.append(fonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "000000")
    run_properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "none")
    run_properties.append(underline)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "24")
    run_properties.append(size)
    size_cs = OxmlElement("w:szCs")
    size_cs.set(qn("w:val"), "24")
    run_properties.append(size_cs)
    run.append(run_properties)
    text = OxmlElement("w:t")
    text.text = display_text
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def build_formal_docx(data: dict[str, Any], template_path: Path, output_path: Path, formal_mode: str) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.text.paragraph import Paragraph

    document = Document(str(template_path))
    run = data["run"]
    project = run["project"]
    formal = require_object(run.get("formal_record", {}), "run.formal_record")

    def formal_value(key: str, placeholder: str) -> str:
        value = formal.get(key)
        if isinstance(value, list):
            text = "、".join(str(item).strip() for item in value if str(item).strip())
        else:
            text = str(value or "").strip()
        return text or (placeholder if formal_mode == "draft" else "")

    replacements = {
        "【项目名称】": project["name"],
        "【项目简称】": project["short_name"],
        "【查验事项】": project["matter"],
        "【核查期间】": project["period"],
        "【查询时间】": formal_value("query_date", "【待填写查询时间】"),
        "【查询地点】": formal_value("query_location", "【待填写查询地点】"),
        "【查询人】": formal_value("query_people", "【待填写查询人】"),
    }
    for paragraph in iter_paragraphs(document):
        replace_in_paragraph(paragraph, replacements)

    marker_row_found = False
    for table in document.tables:
        for row in list(table.rows):
            if "【SUBJECT_ROWS】" not in "".join(cell.text for cell in row.cells):
                continue
            marker_row_found = True
            for subject in data["subjects"]:
                new_row = table.add_row()
                row._tr.addprevious(new_row._tr)
                if len(new_row.cells) < 3:
                    raise ValidationError("DOCX template subject table must contain at least three columns.")
                identifier = formal_identifier_text(subject, formal_mode)
                set_cell_text(new_row.cells[0], subject["role"])
                set_cell_text(new_row.cells[1], subject["name"])
                set_cell_text(new_row.cells[2], identifier)
            table._tbl.remove(row._tr)
    if not marker_row_found:
        raise ValidationError("DOCX template is missing 【SUBJECT_ROWS】 marker row.")

    query_marker = None
    for paragraph in document.paragraphs:
        if "【QUERY_RESULT_ITEMS】" in paragraph.text:
            query_marker = paragraph
            break
    if query_marker is None:
        raise ValidationError("DOCX template is missing 【QUERY_RESULT_ITEMS】 marker paragraph.")

    subject_by_id = {item["subject_id"]: item for item in data["subjects"]}
    for index, query in enumerate(data["queries"], start=1):
        new_xml = copy.deepcopy(query_marker._p)
        query_marker._p.addprevious(new_xml)
        paragraph = Paragraph(new_xml, query_marker._parent)
        clear_paragraph_numbering(paragraph)
        subject = subject_by_id[query["subject_id"]]
        query_terms, filters = query_display_descriptions(
            str(data.get("schema_version") or ""),
            subject,
            query,
        )
        prefix, url, suffix = concise_formal_query_parts(
            index,
            query,
            subject["name"],
            query_terms=query_terms,
            filters=filters,
        )
        clear_paragraph_content(paragraph)
        prefix_run = paragraph.add_run(prefix)
        set_run_font(prefix_run, 12)
        append_external_hyperlink(paragraph, url, breakable_url_display(url))
        suffix_run = paragraph.add_run(suffix)
        set_run_font(suffix_run, 12)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.first_line_indent = None
        paragraph.paragraph_format.left_indent = None
        paragraph.paragraph_format.right_indent = None
        paragraph.paragraph_format.line_spacing = 1.5
        enable_latin_word_wrap(paragraph)
    if formal_mode == "draft" and str(data.get("schema_version") or "") == "1.2":
        execution_by_scope = {
            item["scope_item_id"]: item for item in data.get("query_executions", [])
        }
        queried_scope = {item["scope_item_id"] for item in data["queries"]}
        next_index = len(data["queries"]) + 1
        for scope_item in data["query_scope"]["items"]:
            if scope_item["scope_item_id"] in queried_scope:
                continue
            new_xml = copy.deepcopy(query_marker._p)
            query_marker._p.addprevious(new_xml)
            paragraph = Paragraph(new_xml, query_marker._parent)
            clear_paragraph_numbering(paragraph)
            clear_paragraph_content(paragraph)
            text = draft_incomplete_scope_text(
                next_index,
                scope_item,
                execution_by_scope.get(scope_item["scope_item_id"]),
                subject_by_id[scope_item["subject_id"]]["name"],
            )
            text_run = paragraph.add_run(text)
            set_run_font(text_run, 12)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            paragraph.paragraph_format.first_line_indent = None
            paragraph.paragraph_format.left_indent = None
            paragraph.paragraph_format.right_indent = None
            paragraph.paragraph_format.line_spacing = 1.5
            next_index += 1
    query_marker._element.getparent().remove(query_marker._element)

    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.title = f"网络查询记录-{project['short_name']}"
    document.core_properties.subject = ""
    document.core_properties.keywords = ""
    document.core_properties.comments = ""
    document.core_properties.category = ""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))


def resolve_template_path(requested: Path | None, formal_mode: str) -> Path | None:
    if formal_mode == "none":
        return None
    if requested is not None:
        if not requested.is_file():
            raise FileNotFoundError(f"Requested DOCX template not found: {requested}")
        template_check(requested)
        return requested
    if BUNDLED_TEMPLATE.is_file():
        template_check(BUNDLED_TEMPLATE)
        return BUNDLED_TEMPLATE
    if formal_mode == "final":
        raise FileNotFoundError(
            "No readable DOCX template is available; final mode cannot continue."
        )
    print(
        "Warning: no readable DOCX template is available; generated Markdown only.",
        file=sys.stderr,
    )
    return None


def build_outputs(
    data: dict[str, Any],
    *,
    workpaper_root: Path,
    output_dir: Path,
    template_docx: Path | None,
    formal_mode: str,
    overwrite: bool,
    layout: str = "two-layer",
) -> list[Path]:
    root = workpaper_root.resolve()
    resolved_output_dir = output_dir.resolve()
    try:
        resolved_output_dir.relative_to(root)
    except ValueError as exc:
        raise ValidationError("output_dir must be inside workpaper_root.") from exc
    if layout not in {"two-layer", "flat"}:
        raise ValidationError("layout must be two-layer or flat.")
    validate_run(data, workpaper_root=workpaper_root, formal_mode=formal_mode)
    resolved_template = resolve_template_path(template_docx, formal_mode)
    formal_enabled = formal_mode != "none" and resolved_template is not None

    queries = data["queries"]
    date_token = "未查询"
    if queries:
        latest = max(require_timestamp(item["query_time"], "query.query_time") for item in queries)
        date_token = latest.astimezone(ZoneInfo(data["run"]["timezone"])).strftime("%Y%m%d")
    name_counts: dict[str, int] = {}
    for subject in data["subjects"]:
        name_counts[subject["name"]] = name_counts.get(subject["name"], 0) + 1

    internal_output_dir = output_dir / "01-内部底稿" if layout == "two-layer" else output_dir
    formal_output_dir = output_dir / "02-正式记录" if layout == "two-layer" else output_dir
    final_paths: list[Path] = []
    subject_queries: dict[str, list[dict[str, Any]]] = {item["subject_id"]: [] for item in data["subjects"]}
    for query in queries:
        subject_queries[query["subject_id"]].append(query)
    for subject in data["subjects"]:
        suffix = f"-{subject['subject_id']}" if name_counts[subject["name"]] > 1 else ""
        filename = f"{safe_filename(subject['name'])}网络核查底稿-{date_token}{suffix}.md"
        final_paths.append(internal_output_dir / filename)
    if formal_enabled:
        project_name = safe_filename(data["run"]["project"]["short_name"])
        draft_suffix = "_草稿" if formal_mode == "draft" else ""
        final_paths.append(
            formal_output_dir / f"网络查询记录-{project_name}-{date_token}{draft_suffix}.docx"
        )

    internal_output_dir.mkdir(parents=True, exist_ok=True)
    if formal_enabled:
        formal_output_dir.mkdir(parents=True, exist_ok=True)
    safe_destinations = [
        safe_output_destination(path, root, "output file") for path in final_paths
    ]
    existing = [path for path in safe_destinations if path.exists()]
    if existing and not overwrite:
        raise FileExistsError("Output file(s) already exist: " + ", ".join(str(path) for path in existing))

    with tempfile.TemporaryDirectory(prefix="network-workpaper-build-") as temp_text:
        temp_dir = Path(temp_text)
        staged: list[tuple[Path, Path]] = []
        path_index = 0
        for subject in data["subjects"]:
            destination = final_paths[path_index]
            path_index += 1
            staged_path = temp_dir / destination.name
            staged_path.write_text(
                build_internal_markdown(
                    data,
                    subject,
                    subject_queries[subject["subject_id"]],
                    workpaper_root=workpaper_root,
                    markdown_dir=internal_output_dir,
                ),
                encoding="utf-8",
            )
            staged.append((staged_path, destination))
        if formal_enabled:
            destination = final_paths[path_index]
            staged_path = temp_dir / destination.name
            assert resolved_template is not None
            build_formal_docx(data, resolved_template, staged_path, formal_mode)
            staged.append((staged_path, destination))
        for (staged_path, _), destination in zip(staged, safe_destinations):
            atomic_copy_file(staged_path, destination)
    return final_paths


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Prepare, validate, and build legal network workpapers.")
    subparsers = parser.add_subparsers(dest="command", required=True)

    preview = subparsers.add_parser(
        "scope-preview",
        help="Print the confirmed subject/issue/site/condition matrix before browsing.",
    )
    preview.add_argument("run_file", type=Path)

    doctor = subparsers.add_parser("doctor", help="Check portable runtime capabilities.")
    doctor.add_argument(
        "--browser",
        choices=("available", "unavailable", "unknown"),
        default="unknown",
        help="Browser capability confirmed by the host Agent or user.",
    )

    completion = subparsers.add_parser(
        "completion-preview",
        help="Print schema 1.2 execution coverage, blockers, and final readiness.",
    )
    completion.add_argument("run_file", type=Path)
    completion.add_argument("--workpaper-root", type=Path, required=True)

    prepare = subparsers.add_parser("prepare", help="Create a run skeleton without network access.")
    prepare.add_argument("input_file", type=Path)
    prepare.add_argument("--output", type=Path, required=True)
    prepare.add_argument("--overwrite", action="store_true")

    validate = subparsers.add_parser("validate", help="Validate a completed run file.")
    validate.add_argument("run_file", type=Path)
    validate.add_argument("--workpaper-root", type=Path, default=None)
    validate.add_argument("--formal-mode", choices=("none", "draft", "final"), default="draft")

    template = subparsers.add_parser(
        "template-check", help="Check DOCX markers, page field and privacy safety."
    )
    template.add_argument("template_docx", type=Path)

    audit = subparsers.add_parser(
        "artifact-audit",
        help="Scan one generated layer or one explicit artifact for sensitive values and local paths.",
    )
    audit.add_argument(
        "artifact_path",
        type=Path,
        help="01-内部底稿, 02-正式记录, or one explicit artifact file; never a project root.",
    )
    audit.add_argument(
        "--run-file",
        type=Path,
        default=None,
        help="Validated run context required to allow its exact company code in the formal identity cell.",
    )

    build = subparsers.add_parser("build", help="Generate internal Markdown and optional formal DOCX.")
    build.add_argument("run_file", type=Path)
    build.add_argument("--workpaper-root", type=Path, required=True)
    build.add_argument("--output-dir", type=Path, required=True)
    build.add_argument("--template-docx", type=Path, default=None)
    build.add_argument("--formal-mode", choices=("none", "draft", "final"), default="draft")
    build.add_argument("--overwrite", action="store_true")
    build.add_argument("--layout", choices=("two-layer", "flat"), default="two-layer")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    try:
        if args.command == "scope-preview":
            preview = scope_preview(load_json(args.run_file))
            print(json.dumps(preview, ensure_ascii=False, indent=2))
        elif args.command == "doctor":
            report = doctor_report(args.browser)
            print(json.dumps(report, ensure_ascii=False, indent=2))
            if not report["python"]["ok"] or not all(report["dependencies"].values()):
                return 2
        elif args.command == "completion-preview":
            preview = completion_preview(
                load_json(args.run_file), workpaper_root=args.workpaper_root
            )
            print(json.dumps(preview, ensure_ascii=False, indent=2))
        elif args.command == "prepare":
            prepared = prepare_run(load_json(args.input_file))
            write_json(args.output, prepared, overwrite=args.overwrite)
            print(args.output)
        elif args.command == "validate":
            validate_run(
                load_json(args.run_file),
                workpaper_root=args.workpaper_root,
                formal_mode=args.formal_mode,
            )
            print("Network verification run is valid.")
        elif args.command == "template-check":
            print(json.dumps(template_check(args.template_docx), ensure_ascii=False, indent=2))
        elif args.command == "artifact-audit":
            run_data = load_json(args.run_file) if args.run_file is not None else None
            report = artifact_audit(args.artifact_path, run_data=run_data)
            print(json.dumps(report, ensure_ascii=False, indent=2))
            if not report["ok"]:
                return 2
        else:
            paths = build_outputs(
                load_json(args.run_file),
                workpaper_root=args.workpaper_root,
                output_dir=args.output_dir,
                template_docx=args.template_docx,
                formal_mode=args.formal_mode,
                overwrite=args.overwrite,
                layout=args.layout,
            )
            for path in paths:
                print(path)
            pending = manual_identifier_fields_pending(load_json(args.run_file))
            if pending:
                print("manual_identifier_fields_pending: " + ", ".join(pending))
    except (FileExistsError, FileNotFoundError, OSError, ValidationError, json.JSONDecodeError) as exc:
        print(str(exc), file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
